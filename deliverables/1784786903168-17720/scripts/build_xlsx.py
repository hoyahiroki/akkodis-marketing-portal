#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AKKODiSマーケティングポータル リニューアル提案書（Excel）生成スクリプト
- サンプル形式（フルサイズのスクショに注釈を重ね、大きく見せる）に倣う
- AsIs（注釈付き実スクショ・大判）を左、ToBe（wf_v5・大判）を右に横並び配置
- 判読性最優先：画像は大きく、行高・列幅は画像サイズに合わせて調整
"""
import os
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.properties import WorksheetProperties, PageSetupProperties
from PIL import Image as PILImage

ROOT = "/workspace/akkodis-marketing-portal"
TASK_DIR = os.path.join(ROOT, "deliverables/1784786903168-17720")
ASSET_DIR = os.path.join(TASK_DIR, "xlsx_assets")
OUT_PATH = os.path.join(TASK_DIR, "AKKODiSマーケティングポータル_リニューアル提案書.xlsx")

# ---------------------------------------------------------------------------
# ブランドトークン（akkodis-brand-core）
# ---------------------------------------------------------------------------
NAVY = "001F33"
YELLOW = "FFB81C"
CYAN = "00FFFF"
WHITE = "FFFFFF"
NAVY_80 = "334B5C"
NAVY_40 = "999AAD"
NAVY_20 = "CCC8D6"
LIGHTBG = "F7F7F5"
BLACK = "000000"
RED_NOTE = "C41E3A"  # 注意喚起用（ブランド外色は使わず、強調は太字＋NAVY/YELLOWで表現。差戻し等の注記のみ最小限）

FONT_NAME = "Meiryo UI"

ROW_H_PT = 15          # 既定行高（pt）。15pt = 20px（96dpi）として画像との整合計算に用いる
ROW_H_PX = 20
COL_W = 15             # 既定列幅（文字単位）。おおよそ110pxに相当
COL_PX = 110

IMG_TARGET_W = 1200    # 事前リサイズ済み画像の幅(px)。xlsx_assets 生成時に統一済み


def F(size=10, bold=False, color=BLACK, italic=False, name=FONT_NAME):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


THIN = Side(style="thin", color=NAVY_40)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
WRAP_TOP_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_rows_height(ws, r0, r1, pt=ROW_H_PT):
    for r in range(r0, r1 + 1):
        ws.row_dimensions[r].height = pt


def set_col_widths(ws, first_col, last_col, width=COL_W):
    for c in range(first_col, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = width


def confidential_tag(ws, row, col):
    c = ws.cell(row=row, column=col, value="社外秘")
    c.font = F(size=9, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")


def page_header_bar(ws, row, first_col, last_col, title, sheet_no, sheet_total=8):
    """全シート共通のヘッダーバー（濃いブルー地・白文字）"""
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col - 2)
    c = ws.cell(row=row, column=first_col, value=f"AKKODiSマーケティングポータル リニューアル提案　|　{title}")
    c.font = F(size=15, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(first_col, last_col - 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=row, start_column=last_col - 1, end_row=row, end_column=last_col)
    c2 = ws.cell(row=row, column=last_col - 1, value=f"社外秘　{sheet_no}/{sheet_total}")
    c2.font = F(size=10, bold=True, color=YELLOW)
    c2.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    c2.fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[row].height = 30


def section_title(ws, row, first_col, last_col, text, fill=NAVY, color=WHITE, size=12, height=24):
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=first_col, value=text)
    c.font = F(size=size, bold=True, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(first_col, last_col + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)
    ws.row_dimensions[row].height = height
    return row + 1


def note_box(ws, row, first_col, last_col, text, height=40, fill=LIGHTBG, bold=False, color=BLACK, size=10,
             italic=False):
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=first_col, value=text)
    c.font = F(size=size, bold=bold, color=color, italic=italic)
    c.alignment = WRAP_TOP
    for col in range(first_col, last_col + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=row, column=col).border = BORDER_ALL
    ws.row_dimensions[row].height = height
    return row + 1


def add_image_at(ws, path, anchor_col, anchor_row, width_px=None, height_px=None):
    w, h = PILImage.open(path).size
    img = XLImage(path)
    img.width = width_px or w
    img.height = height_px or h
    ws.add_image(img, f"{get_column_letter(anchor_col)}{anchor_row}")
    return img.width, img.height


# ---------------------------------------------------------------------------
# 区分（ステータス）: 「全22課題を解消」という一括表現を廃止し、課題ごとに区分を明示する。
# ---------------------------------------------------------------------------
STATUS_RESOLVED = "解消"
STATUS_KEEP = "社長指示により現状維持"
STATUS_OUT = "対象外"
STATUS_GOOD = "Good（維持・横展開）"
STATUS_CHECK = "要現物確認"

STATUS_COLOR = {
    STATUS_RESOLVED: ("E7F2EA", NAVY_80),
    STATUS_KEEP: ("EFEFEF", NAVY_80),
    STATUS_OUT: ("EFEFEF", NAVY_40),
    STATUS_GOOD: ("FFF4E0", NAVY_80),
    STATUS_CHECK: ("FDEBEC", RED_NOTE),
}


# ---------------------------------------------------------------------------
# 汎用: 課題→区分→改善→期待効果テーブル
# ---------------------------------------------------------------------------
def issue_table(ws, start_row, first_col, id_span, status_span, issue_span, improve_span, effect_span, rows,
                 headers=("課題ID", "区分", "課題（AsIs）", "改善内容（ToBe）", "期待効果")):
    r = start_row
    c0 = first_col
    c1 = c0 + id_span - 1
    c_s1 = c1 + 1
    c_s2 = c_s1 + status_span - 1
    c2 = c_s2 + issue_span
    c3 = c2 + improve_span
    c4 = c3 + effect_span
    spans = [(c0, c1), (c_s1, c_s2), (c_s2 + 1, c2), (c2 + 1, c3), (c3 + 1, c4)]

    for (a, b), h in zip(spans, headers):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=h)
        cell.font = F(size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = CENTER
        cell.border = BORDER_ALL
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[r].height = 22
    r += 1

    for i, row_vals in enumerate(rows):
        shade = LIGHTBG if i % 2 == 0 else WHITE
        status = row_vals[1]
        status_fill, status_font_color = STATUS_COLOR.get(status, (shade, BLACK))
        for (a, b), val in zip(spans, row_vals[:5]):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            cell = ws.cell(row=r, column=a, value=val)
            is_status_col = a == c_s1
            cell.font = F(size=10, bold=(a == c0 or is_status_col), color=(status_font_color if is_status_col else BLACK))
            cell.alignment = WRAP_TOP_CENTER if a in (c0, c_s1) else WRAP_TOP
            fill_color = status_fill if is_status_col else shade
            for col in range(a, b + 1):
                ws.cell(row=r, column=col).border = BORDER_ALL
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill_color)
        ws.row_dimensions[r].height = 92
        r += 1
    return r


# ===========================================================================
# Workbook
# ===========================================================================
wb = Workbook()
wb.remove(wb.active)

TOTAL_SHEETS = 9

# ---------------------------------------------------------------------------
# Sheet 1: 表紙
# ---------------------------------------------------------------------------
ws = wb.create_sheet("表紙")
ws.sheet_view.showGridLines = False
set_col_widths(ws, 1, 14, width=10)

logo_path = os.path.join(ASSET_DIR, "AKKODIS_Logo_POS_RGB.png")
lw, lh = PILImage.open(logo_path).size
scale = 260 / lw
add_image_at(ws, logo_path, 2, 3, width_px=260, height_px=round(lh * scale))
set_rows_height(ws, 1, 12, pt=22)

ws.merge_cells("B12:M12")
c = ws.cell(row=12, column=2, value="社外秘（Confidential）")
c.font = F(size=11, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=NAVY)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[12].height = 26

ws.merge_cells("B15:N17")
c = ws.cell(row=15, column=2, value="AKKODiSマーケティングポータル\nリニューアル提案")
c.font = F(size=28, bold=True, color=NAVY)
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
set_rows_height(ws, 15, 17, pt=26)

ws.merge_cells("B19:N21")
c = ws.cell(
    row=19, column=2,
    value="現状課題22件の可視化と、SharePoint標準実装によるToBe（ワイヤーフレーム v5・承認済）\n"
          "AI-1（ブランドセルフチェッカー）／AI-2（マーケ特化AIアシスタント）実装をマスト要件として統合",
)
c.font = F(size=13, color=NAVY_80)
c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
set_rows_height(ws, 19, 21, pt=22)

info_rows = [
    ("対象", "AKKODiS マーケティングポータル（SharePoint Online イントラサイト）"),
    ("位置づけ", "第1フェーズ現状分析（docs/01）→ ToBe WF v5（社長4回差戻し承認）を踏まえた提案書化フェーズ"),
    ("作成日", "2026-07-30"),
    ("作成", "AI Corporation デザイン部署（designer-coder）／マーケ部署の一次情報に基づく"),
    ("版数", "v1.0（WF v5 確定版に基づく）"),
]
r = 24
for label, val in info_rows:
    ws.cell(row=r, column=2, value=label).font = F(size=10, bold=True, color=WHITE)
    ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=NAVY_80)
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=13)
    vcell = ws.cell(row=r, column=3, value=val)
    vcell.font = F(size=10)
    vcell.alignment = WRAP_TOP
    for col in range(2, 14):
        ws.cell(row=r, column=col).border = BORDER_ALL
    ws.row_dimensions[r].height = 22
    r += 1

ws.merge_cells(f"B{r+2}:N{r+6}")
note = ws.cell(
    row=r + 2, column=2,
    value="本書はイントラ情報（社外秘を含む）を扱うため、成果物はリポジトリ内のみに保存し外部への持ち出し・転載は行っていません。\n"
          "事実は docs/01_current-site-analysis.md および差戻し回答書_v2.md〜_v4.md・sharepoint-verify-v3.md〜v5.md（一次情報・出典URL付き）に忠実に記載し、"
          "創作は行っていません。文言未確定の箇所は「（マーケ部確定待ち）」と明記しています。根拠資料の全リストは「AI-1,2実装と確認事項」シート末尾を参照。",
)
note.font = F(size=9.5, italic=True, color=NAVY_80)
note.alignment = WRAP_TOP
set_rows_height(ws, r + 2, r + 6, pt=20)


# ---------------------------------------------------------------------------
# Sheet 2: エグゼクティブサマリー
# ---------------------------------------------------------------------------
ws = wb.create_sheet("エグゼクティブサマリー")
ws.sheet_view.showGridLines = False
LAST_COL_2 = 16
set_col_widths(ws, 1, LAST_COL_2, width=12)
page_header_bar(ws, 1, 1, LAST_COL_2, "エグゼクティブサマリー", 2, TOTAL_SHEETS)

r = section_title(ws, 3, 1, LAST_COL_2, "1. 現状課題 22件の内訳（docs/01_current-site-analysis.md）")
r += 1
cat_rows = [
    ("T", "トップ（ホーム）", "4", "T-1〜T-4"),
    ("B", "ブランドライブラリ（最重要）", "7", "B-1〜B-7"),
    ("C", "クライアント向け", "3", "C-1〜C-3"),
    ("CA", "キャンディデート向け", "2", "CA-1〜CA-2"),
    ("P", "グローバルブランドパートナーシップ", "2", "P-1〜P-2"),
    ("G", "全ページ横断", "4", "G-1〜G-4"),
    ("合計", "", "22", ""),
]
headers = ["区分", "ページ／範囲", "件数", "課題ID"]
spans = [(1, 2), (3, 7), (8, 9), (10, 16)]
for (a, b), h in zip(spans, headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 20
r += 1
for i, row_vals in enumerate(cat_rows):
    is_total = row_vals[0] == "合計"
    shade = NAVY_20 if is_total else (LIGHTBG if i % 2 == 0 else WHITE)
    for (a, b), val in zip(spans, row_vals):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        cell.font = F(bold=is_total)
        cell.alignment = CENTER if a != 3 else CENTER
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=shade)
    ws.row_dimensions[r].height = 20
    r += 1

r += 1
r = section_title(ws, r, 1, LAST_COL_2, "1-2. 課題別ステータス内訳（「全22課題を解消」という一括表現は用いない）", fill=NAVY_80, size=11)
r += 1
status_breakdown = [
    (STATUS_RESOLVED, "14", "設計変更により改善するもの（T-3,T-4／B-1〜B-4,B-6,B-7／C-1,C-2／P-2／G-1〜G-3）"),
    (STATUS_KEEP, "2", "社長指示により設計を変更せず現状維持（T-2＝組織体制の配置／CA-1＝各チーム連絡先の現行踏襲）"),
    (STATUS_OUT, "2", "社長指示・確認により今回のスコープ対象外（T-1＝モバイル考慮外／G-4＝権限で非表示のため問題なし）"),
    (STATUS_GOOD, "2", "既に良好なため他ページへ横展開（C-3・CA-2）"),
    (STATUS_CHECK, "2", "確認の結果、修正不要または現物確認が必要（B-5＝ラベル付きと判明し修正不要／P-1＝ASP画像枠の要現物確認）"),
    ("合計", "22", ""),
]
sb_headers = ["区分", "件数", "内容"]
sb_spans = [(1, 4), (5, 6), (7, 16)]
for (a, b), h in zip(sb_spans, sb_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 20
r += 1
for status, count, desc in status_breakdown:
    is_total = status == "合計"
    status_fill, status_font_color = STATUS_COLOR.get(status, (NAVY_20 if is_total else WHITE, BLACK))
    for (a, b), val in zip(sb_spans, (status, count, desc)):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        cell.font = F(bold=(a in (1, 5) or is_total), color=(status_font_color if a == 1 and not is_total else BLACK))
        cell.alignment = CENTER if a in (1, 5) else WRAP_TOP
        fill_color = (NAVY_20 if is_total else (status_fill if a == 1 else WHITE))
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill_color)
    ws.row_dimensions[r].height = 26
    r += 1

r += 1
r = section_title(ws, r, 1, LAST_COL_2, "2. 最重要課題トップ3（docs/01 §4-1・優先着手）")
top3 = [
    ("B-1", "ブランドライブラリの情報過多・ページ内ナビ欠如",
     "インパクト高／コスト中〜高。直近30日で最も使われる実働ページ（一意閲覧者319名）かつユーザー指摘「ガイドラインが増えてわかりにくい」に完全一致。改善効果が最も広く届く。"),
    ("G-2", "ポータル内の検索性の弱さ",
     "インパクト高／コスト中〜高。「見つからない→メール検索」という核心症状の根本原因。AI-2（マーケ特化AIアシスタント）導入の主戦場。"),
    ("CA-1", "資料が「各チームに直接連絡」で完結し入手不可",
     "インパクト高／コスト低。ユーザー指摘「見つからずメール検索」をポータル側が誘発している典型例（社長確認の結果、現状維持を採用＝v3以降）。"),
]
for cid, title, desc in top3:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c1 = ws.cell(row=r, column=1, value=cid)
    c1.font = F(bold=True, color=WHITE)
    c1.fill = PatternFill("solid", fgColor=NAVY_80)
    c1.alignment = CENTER
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=16)
    c2 = ws.cell(row=r, column=3, value=f"{title}\n{desc}")
    c2.font = F(size=10)
    c2.alignment = WRAP_TOP
    for col in range(1, 17):
        ws.cell(row=r, column=col).border = BORDER_ALL
    ws.row_dimensions[r].height = 46
    r += 1

r += 1
r = section_title(ws, r, 1, LAST_COL_2, "3. 改善方針")
policies = [
    "改善対象14課題を SharePoint 標準Web パーツ（クイックリンク／テキスト／Call to action／埋め込み 等）で実装し、追加ライセンス無しで実現できる範囲（L1）を中心に構築する。現状維持・対象外・Good維持・要現物確認の8課題は設計変更を行わず、上記「1-2. 課題別ステータス内訳」のとおり状態を明示する。",
    "AI-1（ブランドセルフチェッカー：Copilot Studio）／AI-2（マーケ特化AIアシスタント：SharePoint agent）の実装をマスト要件とする（社長指示・クイックウィンの段階投入は行わない）。",
    "改善対象14課題＋AI-1,2を1つのリリースとして同日カットオーバー公開する（F0〜F4のフェーズ計画。詳細は「フェーズ計画」シート）。",
    "AI-1,2実装の前提となる M365 Copilot／Copilot Studio のライセンス・管理者許可の確認（F0：テナント・ライセンス棚卸し）が最大のクリティカルパス。F0でライセンス・権限が確保できない場合は、同日公開の前提が崩れるため公開延期が原則（Must要件を外した部分公開は行わない）。",
]
for p in policies:
    r = note_box(ws, r, 1, LAST_COL_2, "・" + p, height=32)

r += 1
r = section_title(ws, r, 1, LAST_COL_2, "4. 全体像（ToBe WF v5・5ページ）")
r += 1
thumb_row = r
thumb_cols = [2, 5, 8, 11, 14]
labels = ["① トップ", "② ブランドライブラリ", "③ クライアント向け", "④ キャンディデート向け", "⑤ パートナーシップ"]
files = ["thumb_top.png", "thumb_brand.png", "thumb_client.png", "thumb_candidate.png", "thumb_partner.png"]
max_h = 0
for col, fname in zip(thumb_cols, files):
    p = os.path.join(ASSET_DIR, fname)
    w, h = PILImage.open(p).size
    add_image_at(ws, p, col, thumb_row)
    max_h = max(max_h, h)
rows_needed = math.ceil(max_h / ROW_H_PX) + 2
set_rows_height(ws, thumb_row, thumb_row + rows_needed)
label_row = thumb_row + rows_needed + 1
for col, lab in zip(thumb_cols, labels):
    ws.merge_cells(start_row=label_row, start_column=col, end_row=label_row, end_column=col + 2)
    c = ws.cell(row=label_row, column=col, value=lab)
    c.font = F(size=9, bold=True, color=NAVY)
    c.alignment = Alignment(horizontal="center")
r = label_row + 2
r = note_box(ws, r, 1, LAST_COL_2,
             "※各ページの大判AsIs／ToBe・課題対応の詳細は「課題・改善内容」シートを参照。", height=24)


# ---------------------------------------------------------------------------
# Sheet 3: 課題・改善内容（メインシート）
# ---------------------------------------------------------------------------
ws = wb.create_sheet("課題・改善内容")
ws.sheet_view.showGridLines = False
LAST_COL_3 = 24
ws.column_dimensions[get_column_letter(1)].width = 3
set_col_widths(ws, 2, LAST_COL_3, width=COL_W)
page_header_bar(ws, 1, 1, LAST_COL_3, "課題・改善内容（ページ別 AsIs → ToBe）", 3, TOTAL_SHEETS)
r = note_box(
    ws, 3, 1, LAST_COL_3,
    "各ページ、左に AsIs（現状・注釈付き実スクショ／materials/screenshots より）、右に ToBe（改善後・ToBe_WF_v5_承認用.pdf より）を大判で並べています。"
    "画像内の番号バッジは docs/01_current-site-analysis.md の課題IDに対応します（凡例は各AsIs画像下部）。"
    "各表の「区分」列は、解消／社長指示により現状維持／対象外／Good（維持・横展開）／要現物確認 の5区分で状態を明示します"
    "（内訳は「エグゼクティブサマリー」シート参照。「全22課題を解消」という一括表現は用いません）。",
    height=44, italic=True,
)
r += 1


def build_issue_block(cur_row, page_title, issue_ids_label, asis_file, tobe_file, table_rows, extra_note=None):
    cur_row = section_title(ws, cur_row, 1, LAST_COL_3, f"{page_title}　｜　対象課題: {issue_ids_label}",
                             fill=NAVY, color=WHITE, size=13, height=26)
    ws.merge_cells(start_row=cur_row, start_column=2, end_row=cur_row, end_column=12)
    c = ws.cell(row=cur_row, column=2, value="AsIs（現状・注釈付き実スクショ）")
    c.font = F(bold=True, color=NAVY)
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=cur_row, start_column=14, end_row=cur_row, end_column=24)
    c2 = ws.cell(row=cur_row, column=14, value="ToBe（改善後・ワイヤーフレーム v5 承認版）")
    c2.font = F(bold=True, color=NAVY)
    c2.alignment = Alignment(horizontal="center")
    ws.row_dimensions[cur_row].height = 20
    cur_row += 1

    img_row = cur_row
    asis_p = os.path.join(ASSET_DIR, asis_file)
    tobe_p = os.path.join(ASSET_DIR, tobe_file)
    aw, ah = add_image_at(ws, asis_p, 2, img_row)
    tw, th = add_image_at(ws, tobe_p, 14, img_row)
    max_h = max(ah, th)
    rows_needed = math.ceil(max_h / ROW_H_PX) + 6
    set_rows_height(ws, img_row, img_row + rows_needed)
    cur_row = img_row + rows_needed + 2

    cur_row = issue_table(ws, cur_row, 2, id_span=2, status_span=3, issue_span=5, improve_span=6, effect_span=7,
                          rows=table_rows)
    if extra_note:
        cur_row = note_box(ws, cur_row, 2, LAST_COL_3, extra_note, height=40, fill="FFF4E0")
    cur_row += 2
    return cur_row


# --- 1. トップ（ホーム） ---------------------------------------------------
top_rows = [
    ("T-1", STATUS_OUT, "埋め込みダッシュボードがモバイル非対応。SP版では極小枠＋横スクロールで指標・グラフが実質判読不能。",
     "社長指示によりモバイル対応は今回のWFスコープ外。PC表示を現状維持し、SP分岐表示はWFに描かない（対象外・根拠: sharepoint-feasibility.md R8「端末別コンテンツ出し分け」＝モダン標準機能では不可）。",
     "リニューアルのスコープを明確化。モバイル対応の要否は次フェーズの経営判断として持ち越す（対応漏れではなく意図的な対象外）。"),
    ("T-2", STATUS_KEEP, "「組織体制」が最下部に配置され、来訪者が最初に求める「探す・依頼する」導線より下位に沈む。",
     "社長指示により配置は最下部のまま変更しない（重要度は現状どおり低）。冒頭の目次（クイックリンク）から④組織体制へのジャンプ導線のみ追加。",
     "配置変更のコストをかけずに、目次からワンクリックで到達できるようにし、発見性を最小コストで改善。"),
    ("T-3", STATUS_RESOLVED, "4カテゴリカードとグローバルナビが完全重複。カードに説明文がなく「次に何ができるか」が伝わらない。",
     "各カード直下に1行の説明文欄を追加（クイックリンク Web パーツ・4項目の横並びを想定。列数は画面幅により変動し公式に固定列数の保証はない＝実機プレビューで確認。sharepoint-verify-v4.md V8）。",
     "クリック前にカードの行き先内容が把握でき、期待値形成ができる。（文言はマーケ部確定待ち）"),
    ("T-4", STATUS_RESOLVED, "「業務依頼フォーム」「ブランドレビューの手順」など複数の入口が並び、使い分けが1画面で判別しづらい。",
     "「①AI-1セルフチェック→②手順ガイド→③業務依頼フォーム」を上から下へ全幅で並べる縦積み3ステップ導線に再設計（v4・社長提案）。各ページに共通の窓口案内ブロックを置き「入口を集約」する（業務プロセス自体をフォーム1つに一本化するものではなく、B-3のフォーム→Teams→Salesforce、CA-1の各チーム連絡等の下流プロセスは現行踏襲）。",
     "依頼前にAI-1で気づきを得たうえで正しい手順に沿って依頼でき、ブランドレビューの差戻し率低減が期待できる。"),
]

# --- 2. ブランドライブラリ（最重要） ---------------------------------------
brand_rows = [
    ("B-1", STATUS_RESOLVED, "情報過多・深いネスト・ページ内ナビ（目次/アンカー）欠如。1ページに全ガイドライン・全手続きを積層し目的の情報に辿り着けない。",
     "冒頭にクイックリンクによる目次を新設。ガイドラインライブラリはセクションの「折りたたみ可能にする」機能で格納（既定：折りたたみ）。",
     "直近30日で最も利用されるページ（一意閲覧者319名）の探索コストを削減。ユーザー指摘「わかりにくい」の直接的解消。"),
    ("B-2", STATUS_RESOLVED, "ガイドライン約10本がフラットに列挙。粒度・目的（ロゴ／写真／動画／文字表記／AI画像 等）で束ねられていない。",
     "目的軸4カテゴリ（全社ブランド規定／制作物別／表記・用語／Akkodis Intelligence）のテキストWeb パーツに再構成（新規ライブラリは作らず既存資料へのリンクを整理）。",
     "目的別に探せるようになり、10本を都度スキャンする負担を解消。"),
    ("B-3", STATUS_RESOLVED, "ロゴ掲載フローがフォーム→Teams→Salesforceをまたぐ6ステップを長文で説明。手続き型コンテンツをテキストのみで表現し実行負荷が高い。",
     "①先方ロゴ掲載／②自社ロゴ提供の2ケースを2カラムセクションで並置し番号付きリストで手順を整理。各カラム末尾に同意書テンプレへの直リンク。AI-1セルフチェック導線を冒頭に組込み。",
     "手続きの見通しが良くなり実行負荷を軽減。AI-1の事前セルフチェックにより申請差戻しの減少が期待できる。"),
    ("B-4", STATUS_RESOLVED, "資料の命名規則が不統一（Akkodis/AKKODiSの表記ゆれ、版・言語・形式の混在）。一覧の走査性を下げる。",
     "表示名・メタデータをAKKODiS表記に統一する運用ルールを明記（列設計）。",
     "一覧の走査性が向上し、表記ブレによる混乱・誤選択を防止。"),
    ("B-5", STATUS_CHECK,
     "（当初）ラベルのないアイコンカード3枚に見え、アフォーダンス不明という改善提案を検討していた。",
     "社長確認により、実際はラベル付きのページ内アンカーリンクカードであることが判明→修正不要（架空の改善提案は撤回。WF本体SECTION8の記述と一致）。根拠：社長差戻しコメント（2026-07-24）による確認。",
     "誤認に基づく不要な改修を回避。※「アイコン」セクション新設はB-5とは別トピック（下記注記参照）。"),
    ("B-6", STATUS_RESOLVED, "冒頭見出しの黄色マーカーが不統一（帯見出しとの二重表現）。編集途中のハイライトが残った可能性があり体裁が粗い。",
     "モダンページのテキストハイライト色は自由指定不可（固定パレットのみ）と判明したため黄色マーカーを廃止し、注意喚起は「セクション背景色」に統一。",
     "ページ全体の視覚言語が統一され、体裁の粗さが解消。全ページ共通のトーンに揃う。"),
    ("B-7", STATUS_RESOLVED, "SP版が極端に長尺（1万px超）。モバイルでの回遊が困難。",
     "B-1〜B-6の再構成（目次新設・折りたたみ格納・4カテゴリ整理）によりページの実効的な情報量・全長を圧縮。抜本的なモバイル専用最適化は次フェーズで検討。",
     "情報整理により結果的にSP版の実効的な長さも縮小。抜本対応（レイアウト分割等）は次フェーズに引き継ぐ（改善度合いは限定的である点を正直に記載）。"),
]

# B-5とは別トピック：「アイコン」セクション新設（既存アイコン素材ライブラリへのリンク追加。WF v5 SECTION8）
brand_icon_note = (
    "追加トピック（B-5とは別）：「アイコン」セクションを新設。UI用／資料用／SNS用アイコンの既存素材ライブラリへ"
    "テキストweb パーツでリンクする（新規ライブラリは作成しない）。ページ内ナビ用のアイコンとは別物であり、"
    "旧B-5（ラベルのないアイコンカード）との混同は解消済み。"
)

# --- 3. クライアント向け ----------------------------------------------------
client_rows = [
    ("C-3", STATUS_GOOD, "（Good）冒頭のアンカー型クイックリンク3件が機能し他ページより回遊配慮がある一方、この良い部分が他ページに横展開されていない。",
     "クライアントページの「冒頭に目次（クイックリンク）を置く」パターンを全ページ共通テンプレートとして横展開（横断課題G-1として全ページに適用）。",
     "ページ間で「探し方」を学び直す必要がなくなり、サイト全体の回遊性が向上。"),
    ("C-1", STATUS_RESOLVED, "「お役立ち資料」が30本超のリンク羅列。説明文・サムネ・絞り込みがなく目的資料の発見コストが高い。",
     "「公開資料ライブラリ」をテキストWeb パーツ（主案）でカテゴリ見出し＋リンク箇条書きに再構成し、30本超はセクション折りたたみで格納。副案としてリスト併用（絞り込み優先）も比較のうえ併記（3構成比較は「クライアント資料の2案比較」シート参照）。",
     "主案採用でAI-2（生成AI）が資料リンクを回答根拠にできる。副案／併用案採用ならカテゴリ/言語での絞り込み・並べ替えが可能（トレードオフは同シート参照）。"),
    ("C-2", STATUS_RESOLVED, "リンクに更新日・対象・要約などのメタ情報がなく、最新版か・誰向けかが判別できない。",
     "更新日は「日付を出せない資料がある」事情を尊重し任意項目のまま維持しつつ、テキストパーツ内の見出し・カテゴリ分類で対象を明確化。",
     "全件必須化を強制せず現実的な運用を維持しながら、カテゴリ単位での対象・種別の視認性を確保（改善度合いは限定的である点を正直に記載）。"),
]

# --- 4. キャンディデート向け -------------------------------------------------
candidate_rows = [
    ("CA-2", STATUS_GOOD, "（Good）5ページ中もっとも簡潔で見やすい構成（AKKODiS公式HP／マスター資料／Social Mediaの3セクション）。",
     "他ページ再設計の参照モデルとしてこの3構成を踏襲。Social Mediaはクイックリンクで4項目の横並びを想定（列数は画面幅により変動・実機プレビューで確認）、AKKODiS Peopleはテキストweb パーツに変更（v4で対応済み）。",
     "改修コストを抑えつつ、他ページとの体験の一貫性を確保する参照モデルとして機能する。"),
    ("CA-1", STATUS_KEEP, "マスター資料が「各チームに直接連絡してください」で完結し、資料そのものが入手できない。ユーザー指摘「メール検索」を誘発する典型例。",
     "社長が担当者に理由を確認中の結果を踏まえ、現行の文言（各チーム連絡先）をテキストWeb パーツでそのまま維持する設計を採用（社長指示により現状維持。理由がある可能性を尊重し直リンク化しない）。",
     "確認結果に基づく意図的な現状維持。直リンク化が可能と判明した場合は再設計の余地を残す（確定案ではなく暫定案）。"),
]

# --- 5. グローバルブランドパートナーシップ -----------------------------------
partner_rows = [
    ("P-1", STATUS_CHECK, "Akkodis ASP Teamのみ余分な（壊れた）画像枠がプレースホルダー表示。他2チームには存在しない画像枠が体裁を崩している（要現物確認）。",
     "画像枠の是正（差し替えまたは除去）。同日リリースに含める低工数項目として扱う（クイックウィン＝段階先行公開ではなく、他の改善と合わせて同日カットオーバーで実施）。実装フェーズで現物確認のうえ対応。",
     "3パートナーの体裁が揃い、プロフェッショナルな見た目を回復する。"),
    ("P-2", STATUS_RESOLVED, "長尺だがページ内ナビがない。クライアントページのアンカー導線と不統一。",
     "ページ内目次（クイックリンク Web パーツ）を新設。3パートナーのサムネイルカードは目次と役割が重複するため削除し、目次に一本化。",
     "クライアントページと同様の回遊導線を確保し、横断課題G-1（ページ間の構成不統一）の解消に寄与（改善度合いは限定的である点を正直に記載）。"),
]

r = build_issue_block(r, "1. トップ（ホーム）", "T-1〜T-4", "asis_top.png", "tobe_top.png", top_rows)
r = build_issue_block(r, "2. ブランドライブラリ（最重要）", "B-1〜B-7", "asis_brand.png", "tobe_brand.png", brand_rows,
                       extra_note=brand_icon_note)
r = build_issue_block(r, "3. クライアント向け", "C-1〜C-3", "asis_client.png", "tobe_client.png", client_rows)
r = build_issue_block(r, "4. キャンディデート向け", "CA-1〜CA-2", "asis_candidate.png", "tobe_candidate.png", candidate_rows)
r = build_issue_block(r, "5. グローバルブランドパートナーシップ", "P-1〜P-2", "asis_partner.png", "tobe_partner.png", partner_rows)


# ---------------------------------------------------------------------------
# Sheet 4: クライアント資料の2案比較（最重要指定）
# ---------------------------------------------------------------------------
ws = wb.create_sheet("クライアント資料2案比較")
ws.sheet_view.showGridLines = False
LAST_COL_4 = 14
set_col_widths(ws, 1, LAST_COL_4, width=13)
page_header_bar(ws, 1, 1, LAST_COL_4, "クライアント資料：3構成比較（①テキストのみ／②テキスト＋リスト併用／③リストのみ）", 4, TOTAL_SHEETS)
r = note_box(
    ws, 3, 1, LAST_COL_4,
    "根拠: 差戻し回答書_v4.md §1-2、sharepoint-verify-v5.md V14〜V17（learn/support.microsoft.com 一次情報）。"
    "AI-1,2の実装がマスト要件であるため、「AI-2が資料リンクを回答根拠にできるか」を最重要の判断軸に据えています。"
    "「リスト併用」を名乗る以上、①テキストのみ／②テキスト＋リスト併用／③リストのみ の3構成で比較します。",
    height=44, italic=True,
)
r += 1

r = section_title(ws, r, 1, LAST_COL_4, "結論（主案・副案）", size=12)
r = note_box(ws, r, 1, LAST_COL_4,
             "主案＝①テキストのみ（AI-2優先）／ 副案＝②テキスト＋リスト併用（絞り込み優先。ただし二重管理コストを伴う）。"
             "③リストのみはAI-2の回答根拠に一切使えないため不採用。どの案で進めるかは社長のご判断を仰ぐ"
             "（差戻し回答書_v4.md §4「ご承認・ご判断いただきたい点」）。",
             height=40, fill="FFF4E0", bold=True)
r += 1

r = section_title(ws, r, 1, LAST_COL_4, "観点別プロコン比較表（3構成）")
compare_headers = ["観点", "① テキストのみ（主案）", "判定", "② テキスト＋リスト併用（副案）", "判定", "③ リストのみ", "判定"]
spans4 = [(1, 2), (3, 5), (6, 6), (7, 9), (10, 10), (11, 13), (14, 14)]
for (a, b), h in zip(spans4, compare_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(size=9, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 30
r += 1

judge_cols = (6, 10, 14)
compare_rows = [
    ("① AI-2（生成AI）が回答根拠にできるか",
     "○　ページ本文としてSharePoint agentのナレッジソース（ページ）に含まれ、拾われる。", "○",
     "○　同じリンクをテキスト側にも残すため、AI-2はテキスト側を参照して回答根拠にできる。", "○",
     "×　SharePoint agentは公式に「リストのデータを使わない（Agents currently don't use data from Lists）」と明記。回答に出てこない。", "×"),
    ("② カテゴリ／言語の絞り込み・並べ替え",
     "×　手作業の見出し分けのみ。動的な絞り込みはできない。", "×",
     "○　リスト側の列・ビューで絞り込み・並べ替えが可能。", "○",
     "○　列・ビューによる絞り込み・並べ替えが可能（既定機能）。", "○"),
    ("③ 件数増加時の管理のしやすさ／二重管理コスト",
     "△　本文が長くなり編集が煩雑になりやすい。", "△",
     "×　同じリンクをテキストとリストの両方に維持する必要があり、二重管理コスト・更新不整合リスク（片方だけ更新して情報が食い違う）が生じる。", "×",
     "○　列で構造的に管理でき、件数増加に強い。", "○"),
    ("④ 人間のサイト内検索でヒットするか",
     "○　ページ本文としてMicrosoft Searchにインデックスされる。", "○",
     "○　テキスト・リスト双方がインデックス対象。", "○",
     "○　リストアイテムも既定でインデックス対象（V14。ただし詳細設定でNoにされていないことが前提）。", "○"),
    ("⑤ 実装／運用コスト",
     "○　テキストパーツの編集のみで完結。新規リスト作成が不要で低コスト。", "○",
     "△　新規リスト作成に加え、同一情報をテキスト側にも重複記載・維持する手間が増え、3構成中もっとも運用コストが高い。", "△",
     "△　新規リストの作成・列/ビュー設計・カテゴリ運用ルールの整備など運用の手間が増える。", "△"),
]
for row_vals in compare_rows:
    vals = list(row_vals)
    for (a, b), val in zip(spans4, vals):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        judge_col = a in judge_cols
        cell.font = F(bold=judge_col, size=13 if judge_col else 9,
                       color=(NAVY_80 if val == "○" else (RED_NOTE if val == "×" else BLACK)) if judge_col else BLACK)
        cell.alignment = CENTER if judge_col else WRAP_TOP
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=LIGHTBG if a == 1 else WHITE)
    ws.row_dimensions[r].height = 58
    r += 1

r += 1
r = note_box(
    ws, r, 1, LAST_COL_4,
    "★二重管理コストの明記（②案）：②テキスト＋リスト併用は、同じ資料リンクをテキストweb パーツとリストの両方に登録・維持する必要があります。"
    "資料の追加・更新・削除のたびに2箇所を同時に直す運用が求められ、片方の更新を忘れると「テキスト側とリスト側で内容が食い違う」不整合リスクが生じます。"
    "絞り込み機能とのトレードオフとして、この運用負荷を許容できるかが③との分岐点です。",
    height=48, fill="FDEBEC", bold=True,
)

r += 1
r = section_title(ws, r, 1, LAST_COL_4, "メリット／デメリット サマリー")
sum_headers = ["① テキストのみ（主案）", "② テキスト＋リスト併用（副案）", "③ リストのみ"]
sum_spans = [(1, 4), (5, 9), (10, 14)]
for (a, b), h in zip(sum_spans, sum_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    c = ws.cell(row=r, column=a, value=h)
    c.font = F(bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY_80)
    c.alignment = CENTER
ws.row_dimensions[r].height = 20
r += 1
sum_vals = [
    "メリット：AI-2が資料リンクを回答根拠にできる（マスト要件のAI導入と両立）。実装・運用コストが最も低い。\n"
    "デメリット：カテゴリ/言語での動的な絞り込み・並べ替えができない（失うものとして正直に明記）。",
    "メリット：AI-2の回答根拠を維持しつつ、リスト側で絞り込み・並べ替えができる。\n"
    "デメリット：テキストとリストの二重管理コスト・更新不整合リスクを常時抱える。3構成中もっとも運用コストが高い。",
    "メリット：カテゴリ/言語での絞り込み・並べ替え、件数増加時の構造的な管理がしやすい。\n"
    "デメリット：AI-2（生成AI）の回答根拠には一切使われない（公式明記の決定的な制約）。新規リスト作成・運用の手間が増える。",
]
sum_fills = ["E7F2EA", "FFF4E0", "FDEBEC"]
for (a, b), val, fill in zip(sum_spans, sum_vals, sum_fills):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    c = ws.cell(row=r, column=a, value=val)
    c.font = F(size=9.5)
    c.alignment = WRAP_TOP
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill)
ws.row_dimensions[r].height = 84
r += 2

r = note_box(
    ws, r, 1, LAST_COL_4,
    "将来リストがSharePoint agentのナレッジソースとして正式対応された場合（非公式情報では対応の兆しがあるが、本調査時点の一次情報では未確認）は運用を見直す価値がある（差戻し回答書_v4.md §2-3）。",
    height=32, italic=True,
)


# ---------------------------------------------------------------------------
# Sheet 5: 横断課題（G-1〜G-4）
# ---------------------------------------------------------------------------
ws = wb.create_sheet("横断課題")
ws.sheet_view.showGridLines = False
LAST_COL_5 = 14
set_col_widths(ws, 1, LAST_COL_5, width=13)
page_header_bar(ws, 1, 1, LAST_COL_5, "全ページ横断の課題（G-1〜G-4）", 5, TOTAL_SHEETS)
r = note_box(ws, 3, 1, LAST_COL_5,
             "根拠: docs/01_current-site-analysis.md §1-6、差戻し回答書.md【全ページ横断の課題】。",
             height=24, italic=True)
r += 1

g_headers = ["課題ID", "区分", "課題", "根拠", "対応方針"]
g_spans = [(1, 2), (3, 4), (5, 7), (8, 9), (10, 14)]
for (a, b), h in zip(g_spans, g_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 22
r += 1

g_rows = [
    ("G-1", STATUS_RESOLVED, "ページ間で構成が不統一（アンカーリンクの有無、装飾の使い方、密度）。回遊のたびに「探し方」を学び直す必要がある。",
     "client=アンカー有／brand・partner=無、brandの見出しマーカー",
     "全5ページで「目次を冒頭に配置／見出しH2統一／注意喚起はセクション背景色に統一」の共通テンプレートを適用。クライアントページのアンカー導線パターン（C-3・Good）をブランド／パートナーへ横展開。"),
    ("G-2", STATUS_RESOLVED, "検索性が弱い。全ページ左上にSharePoint標準検索「このサイトを検索」はあるが、ページ内の絞り込み・ファセット・タグがない。ユーザー指摘「検索性が悪い／見つからない」の核心。",
     "全ページ共通ヘッダーに標準検索は存在するが、ページ内の絞り込み・ファセットUIは無い",
     "本文への検索ボックス追加・位置変更は仕様上不可（sharepoint-feasibility.md R1「検索ボックスの設置位置」）。ヘッダー既定検索を維持しつつ、AI-2（マーケ特化AIアシスタント／SharePoint agent）を導入し自然言語での案内・出典リンク提示で発見性を補う（マスト要件）。"),
    ("G-3", STATUS_RESOLVED, "窓口・相談導線が分散（業務依頼フォーム／ブランドレビュープロセス／各チームへ直接連絡／Salesforceサービスリクエスト／Teams連絡）。",
     "brand・candidate・top",
     "トップの「マーケへの業務依頼」（AI-1セルフチェック→手順ガイド→業務依頼フォームの縦積み3ステップ）を軸に、各ページに共通の窓口案内ブロックを置き「入口を集約」する（業務プロセス自体の一本化ではない。B-3のフォーム→Teams→Salesforce、CA-1の各チーム連絡は現行維持）。"),
    ("G-4", STATUS_OUT, "ナビに「編集 のナビゲーション」が露出。編集者ビュー由来と見られ、一般ユーザー表示では非表示の可能性が高いが要確認。",
     "全ページのナビ末尾",
     "社長確認により、権限のないユーザーには非表示のため対応不要（対象外）。F0（テナント・ライセンス棚卸し）の運用確認事項として念のため実機確認する。"),
]
for i, row_vals in enumerate(g_rows):
    shade = LIGHTBG if i % 2 == 0 else WHITE
    status = row_vals[1]
    status_fill, status_font_color = STATUS_COLOR.get(status, (shade, BLACK))
    for (a, b), val in zip(g_spans, row_vals):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        is_status_col = a == 3
        cell.font = F(bold=(a in (1, 3)), color=(status_font_color if is_status_col else BLACK))
        cell.alignment = WRAP_TOP_CENTER if a in (1, 3) else WRAP_TOP
        fill_color = status_fill if is_status_col else shade
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill_color)
    ws.row_dimensions[r].height = 90
    r += 1


# ---------------------------------------------------------------------------
# Sheet 6: フェーズ計画
# ---------------------------------------------------------------------------
ws = wb.create_sheet("フェーズ計画")
ws.sheet_view.showGridLines = False
LAST_COL_6 = 19
set_col_widths(ws, 1, LAST_COL_6, width=13)
page_header_bar(ws, 1, 1, LAST_COL_6, "フェーズ計画（改善対象14課題＋AI-1,2 同日カットオーバー公開）", 6, TOTAL_SHEETS)
r = note_box(
    ws, 3, 1, LAST_COL_6,
    "前提: 改善対象14課題＋AI-1,2の実装はマスト要件（現状維持・対象外・Good維持・要現物確認の8課題は設計変更しない。内訳は「エグゼクティブサマリー」シート参照）。"
    "クイックウィン／段階投入の考え方は採らず、改善対象14課題＋AI-1,2を1つのリリースとして同日カットオーバー公開する（差戻し回答書.md／差戻し回答書_v2.md）。"
    "期間はいずれも目安であり、実際の所要はF0の確認結果・体制により変動する。",
    height=44, italic=True,
)
r += 1
r = note_box(
    ws, r, 1, LAST_COL_6,
    "★ クリティカルパス／原則：F0のAIライセンス確定。AI-1,2はCopilot Studio／M365 Copilot等のライセンス・管理者許可に依存しており、"
    "F0でAI-1,2のライセンス・権限が確保できない場合は「同日公開」の前提が崩れるため、公開延期が原則（Must要件を外した部分公開は行わない）。"
    "ライセンス単価はF0で確定・要見積（空欄にせず「F0で確定」と明記する）。IT窓口の特定が最優先。",
    height=48, fill="FFF4E0", bold=True,
)
r += 1

ph_headers = ["フェーズ", "期間目安", "主な作業", "成果物", "完了条件", "承認ゲート（Go/No-Go）", "主担当"]
ph_spans = [(1, 2), (3, 4), (5, 8), (9, 10), (11, 13), (14, 17), (18, 19)]
for (a, b), h in zip(ph_spans, ph_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(size=9, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 30
r += 1

phases = [
    ("F0　テナント・ライセンス棚卸し\n（クリティカルパス）",
     "目安2〜4週間\n（ライセンス確認の実際の所要により変動）",
     "M365 Copilot／Copilot Studioの契約・ライセンス割当確認、DLP/機密ラベル設定確認、SharePoint管理者・IT窓口の特定（社長経由）。ライセンス単価の見積取得。",
     "ライセンス・機能可否一覧、AI-1,2実装可否の確定表、ライセンス単価見積",
     "AI-1,2の実装可否・費用（ライセンス単価）が確定する",
     "Go/No-Go①：AI-1,2のライセンス・権限が確保できるか。確保できない場合は同日公開の前提が崩れるため公開延期が原則（Must要件を外した部分公開は行わない）。",
     "社長→SharePoint管理者・IT窓口\n（マーケティング責任者支援）"),
    ("F1　要件定義・情報設計",
     "目安2〜3週間",
     "全ページのToBe（本WF v5承認版）＋AI-1,2の要件・権限（所有者/管理者の切り分け）・スコープを一括確定。クライアント資料の3構成（①②③）から採用案を確定。",
     "要件定義書、確定WF、権限マトリクス",
     "Must範囲（改善対象14課題＋AI-1,2）を社長承認",
     "Go/No-Go②：要件定義書・確定WF・権限マトリクスを社長承認。",
     "マーケティング責任者／SharePoint管理者・IT窓口"),
    ("F2　構築",
     "目安4〜8週間\n（規模により変動）",
     "SharePoint（L1標準Web パーツ）とAI（Copilot Studio／SharePoint agent）を並行構築。",
     "動作するSharePointページ一式＋AI-1,2エージェント試作",
     "改善対象14課題の実装完了、AI-1,2が動作する",
     "Go/No-Go③：構築物の社内一次レビュー通過（QA着手可否）。",
     "実装担当"),
    ("F3　全体検証・移行",
     "目安2〜3週間",
     "改善対象14課題の解消を検証、現状維持・対象外・Good維持・要現物確認8課題の状態確認、既存コンテンツの移行、権限・セキュリティ確認、AI-2の外部リンク提示（要実機検証項目）を含むアクセス権範囲の実機検証。",
     "検証済みステージング環境、QA報告",
     "QA完了・社長承認",
     "Go/No-Go④：QA完了・受入条件の充足・社長最終承認。",
     "実装担当／情報セキュリティ"),
    ("F4　同日カットオーバー公開",
     "目安1日\n（公開作業自体。事前リハーサル別途）",
     "全ページ＋AI-1,2を1つのリリースとして一斉公開。",
     "本番公開ポータル",
     "全ページ＋AI-1,2が同時に稼働開始",
     "（F3のGo/No-Go④通過後に実施。以降の追加ゲートなし）",
     "マーケティング責任者・SharePoint管理者・実装担当"),
]
for i, row_vals in enumerate(phases):
    shade = "FFF4E0" if i == 0 else (LIGHTBG if i % 2 == 0 else WHITE)
    for (a, b), val in zip(ph_spans, row_vals):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        cell.font = F(size=9.5, bold=(a == 1))
        cell.alignment = WRAP_TOP
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=shade)
    ws.row_dimensions[r].height = 100
    r += 1


# ---------------------------------------------------------------------------
# Sheet 7: KPI
# ---------------------------------------------------------------------------
ws = wb.create_sheet("KPI")
ws.sheet_view.showGridLines = False
LAST_COL_7 = 17
set_col_widths(ws, 1, LAST_COL_7, width=13)
page_header_bar(ws, 1, 1, LAST_COL_7, "KPI（社長が取得可能な指標を主軸に3層で設計）", 7, TOTAL_SHEETS)
r = note_box(
    ws, 3, 1, LAST_COL_7,
    "根拠: 差戻し回答書_v4.md §4（権限検証）、sharepoint-verify-v5.md V16。GAは使えないため、SharePoint標準のSite usage／Search usage reportsで測定する。"
    "Search usage reports（ゼロ件検索等）はサイトコレクション管理者のみ閲覧可のため、社長が自分で取れるSite usageを主軸に組み替えた（旧「空振り率」中心の設計を撤回）。"
    "AI-1とAI-2は取得可否が大きく異なるため指標を分離しています（#7・#8）。",
    height=48, italic=True,
)
r += 1

kpi_headers = ["#", "指標", "測定方法", "評価時点", "目標設定の方法", "取得者"]
kpi_spans = [(1, 1), (2, 5), (6, 9), (10, 12), (13, 15), (16, 17)]

EVAL_TIMING_STD = "ベースライン（公開前）取得→公開後30日→公開後90日で追跡し推移を比較。"


def kpi_group(r, title, fill, rows):
    r = section_title(ws, r, 1, LAST_COL_7, title, fill=fill, color=WHITE, size=12, height=22)
    for (a, b), h in zip(kpi_spans, kpi_headers):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=h)
        cell.font = F(size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY_80)
        cell.alignment = CENTER
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
    ws.row_dimensions[r].height = 26
    r += 1
    for i, row_vals in enumerate(rows):
        shade = LIGHTBG if i % 2 == 0 else WHITE
        for (a, b), val in zip(kpi_spans, row_vals):
            ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            cell = ws.cell(row=r, column=a, value=val)
            cell.font = F(size=9.5, bold=(a == 1))
            cell.alignment = CENTER if a in (1, 16) else WRAP_TOP
            for col in range(a, b + 1):
                ws.cell(row=r, column=col).border = BORDER_ALL
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=shade)
        ws.row_dimensions[r].height = 68
        r += 1
    return r + 1


rows_a = [
    ("1", "主要ページの閲覧数・一意閲覧者数（ブランドライブラリ、キャンディデート等）",
     "Site usage（サイト利用状況）のページ別閲覧数・一意閲覧者数を、刷新前後で比較。",
     EVAL_TIMING_STD, "現状値（ベースライン）比で+X%等の相対目標を関係者と協議のうえ設定（絶対値の事前目標は置かない）。",
     "社長（サイト所有者）"),
    ("2", "人気コンテンツ（どの資料/ページが実際に使われているか）",
     "Site usageの「人気コンテンツ（Popular content）」。刷新で導線を張った資料の閲覧数増加を確認。",
     EVAL_TIMING_STD, "導線を張った資料が人気コンテンツ上位に入ることを目標とする（順位ベース、絶対数値目標は置かない）。",
     "社長（サイト所有者）"),
    ("3", "平均滞在時間（探索コストの目安）",
     "Site usageの平均滞在時間。※注記：滞在時間の短縮＝改善とは限らない（探しやすくなり短時間で用が済むケースと、"
     "探せず離脱するケースの両方があり得るため、閲覧数・人気コンテンツ等の他指標と併せて解釈する）。",
     EVAL_TIMING_STD, "単独の数値目標は設定せず、他指標（#1,#2,#6）とセットで傾向を評価する。",
     "社長（サイト所有者）"),
    ("4", "サイト全体の訪問数・利用者数の推移",
     "Site usageの「Site visits／Unique viewers／Popular Platforms（デバイス傾向）」。",
     EVAL_TIMING_STD, "現状値（ベースライン）比での増加傾向を確認（絶対値の事前目標は置かない）。",
     "社長（サイト所有者）"),
]
rows_b = [
    ("5", "ゼロ件検索・人気検索語（＝「見つからない」の直接的な証拠）",
     "Search usage reports（Site settings > Site collection administration > Microsoft Search > Insights）。"
     "月次でサイトコレクション管理者にExcelエクスポートを依頼する運用を提案（権限付与が可能ならそれが最善）。",
     EVAL_TIMING_STD, "現状値（ベースライン）比でのゼロ件検索率の減少を目標とする。",
     "サイトコレクション管理者\n（依頼ベース）"),
]
rows_c = [
    ("6", "マーケ部への「資料が見つからない」問い合わせ件数",
     "手動集計（現状値を取得→刷新後の減少を確認）。",
     EVAL_TIMING_STD, "現状値（ベースライン）比での件数減少を目標とする。",
     "マーケ部（手動）"),
    ("7", "AI-1の利用状況・ブランドレビュー申請数",
     "Copilot Studioの分析機能、またはLists/Forms等の標準集計機能で件数を取得（実装する仕組みの設計に依存）。",
     EVAL_TIMING_STD, "公開後30日の実績値を踏まえてF3以降に目標値を設定する（事前の絶対値目標は置かない）。",
     "マーケ部（Lists/Forms等の集計）"),
    ("8", "AI-2の利用状況",
     "取得方法・権限・保持期間は要確認（SharePoint agentの利用ログ取得機能について一次情報で確認できておらず、"
     "Lists/Formsで取得できるとは書かない。F3の実機検証で確認する）。",
     "F3実機検証時に確認方法を確定のうえ、以降は他指標と同様の時点で追跡。",
     "測定方法確定後に設定（現時点では未確定）。",
     "要確認（IT窓口経由の可能性）"),
]

r = kpi_group(r, "A. 社長ご自身で取得可能（Site usage・主軸）", NAVY, rows_a)
r = kpi_group(r, "B. 管理者への依頼が必要（Search usage reports）", NAVY_80, rows_b)
r = kpi_group(r, "C. 権限不要の手動指標・AI利用状況（補完）", NAVY_40, rows_c)

r = note_box(
    ws, r, 1, LAST_COL_7,
    "運用: 刷新前に上記の現状値（ベースライン）を取得→目標値を設定→公開後30日・90日で測定。Search usage reportsは過去31日（日次）/12か月（月次）でExcel出力可能。"
    "AI-2の利用状況（#8）は取得可否自体が未確認のため、F3実機検証の結果次第で本表を更新する。",
    height=40, italic=True,
)


# ---------------------------------------------------------------------------
# Sheet 8: AI-1,2 実装と確認事項
# ---------------------------------------------------------------------------
ws = wb.create_sheet("AI-1,2実装と確認事項")
ws.sheet_view.showGridLines = False
LAST_COL_8 = 14
set_col_widths(ws, 1, LAST_COL_8, width=13)
page_header_bar(ws, 1, 1, LAST_COL_8, "AI-1・AI-2 実装方式と確認事項", 8, TOTAL_SHEETS)
r = note_box(
    ws, 3, 1, LAST_COL_8,
    "根拠: AI-1-2実装調査.md（開発部署 インフラ/基盤スペシャリスト・learn/support.microsoft.com一次情報）、差戻し回答書_v2.md／_v3.md／_v4.md。",
    height=24, italic=True,
)
r += 1

r = section_title(ws, r, 1, LAST_COL_8, "実装方式")
impl_headers = ["項目", "AI-1　ブランドセルフチェッカー", "AI-2　マーケ特化AIアシスタント"]
impl_spans = [(1, 2), (3, 8), (9, 14)]
for (a, b), h in zip(impl_spans, impl_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 22
r += 1

impl_rows = [
    ("実装方式", "Copilot Studioのカスタムエージェント（会話内ファイルアップロード＋画像入力分析）。SharePointページには「埋め込み Web パーツ（Embed web part）」でカスタムWebサイトチャネルの埋め込みで設置。",
     "SharePointの「エージェント」機能（SharePoint agent）。サイト所有者権限で作成し「メインエージェント」に設定するとヘッダーのエージェントアイコンから全員が利用可。"),
    ("配置場所", "「ブランドレビュー依頼」「ロゴ掲載申請」の文脈（トップSECTION4・ブランドライブラリSECTION3/6）に埋め込み Web パーツで設置。",
     "サイト全体の「メインエージェント」としてトップページに設定。全ページのヘッダーから利用可能。"),
    ("社長（サイト所有者）でできること",
     "Copilot Studioでのエージェント試作・トピック設計・ナレッジ登録・ファイルアップロード機能ON（テナントにCopilot Studio契約と作成者権限があれば）。",
     "SharePoint agentの作成・ナレッジ選定（対象ライブラリ/ページ）・メインエージェント設定・共有制御。ライセンスさえ揃えば所有者だけで作成可能。"),
    ("管理者に確認すべきこと",
     "①M365 Copilot／Copilot Studioライセンスの契約・割当 ②Copilot Studio Authorsロールの付与 ③DLP/機密ラベルの参照除外設定 ④カスタムWebサイトチャネル公開時の「認証なし」設定の許容可否 ⑤PPTX直接アップロード（実験的機能）の申請要否",
     "①M365 Copilotライセンスの契約・割当 ②DLP/機密ラベルの参照除外設定 ③SharePoint Advanced Managementの利用可否 ④Graph connectors（外部データ取込）の要否"),
    ("回答範囲（AI-2・要実機検証）",
     "－（AI-1は画像入力・ファイル解析が主用途のため対象外）",
     "「ページ本文を根拠に回答できる」ことは公式に確認済み。一方「ページ本文中の外部URLを期待どおり提示できるか」は一次情報に明記がなく、F3で必須の実機検証項目とする（sharepoint-verify-v4.md V12、差戻し回答書_v3.md §「クライアントページのHPに飛ぶリンクを探し当てられるか」）。"),
    ("限界・注意点",
     "色のΔE数値照合やロゴのピクセル単位検出など厳密判定は標準機能では非保証。「一次スクリーニング（気づきの提示）」用途が現実的で、最終承認は人が行う運用を推奨（品質保証プロセスは残す）。PPTX直接アップロードは実験的機能、当面はPDF変換を前段に挟む運用が現実的。",
     "回答は必ずユーザー自身のアクセス権範囲でセキュリティトリミングされる（権限境界は標準装備）。参照元への出典（citation）リンクが標準で付与される。"),
]
for i, row_vals in enumerate(impl_rows):
    shade = LIGHTBG if i % 2 == 0 else WHITE
    for (a, b), val in zip(impl_spans, row_vals):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        cell.font = F(bold=(a == 1))
        cell.alignment = WRAP_TOP_CENTER if a == 1 else WRAP_TOP
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=shade)
    ws.row_dimensions[r].height = 96
    r += 1

r += 1
r = note_box(
    ws, r, 1, LAST_COL_8,
    "★重要（決定的な制約）：AI-2（SharePoint agent）は現行仕様で「リストのデータを使わない」ことが公式FAQに明記されています"
    "（\"Agents currently don't use data from Lists.\"）。クライアント公開資料等でAI-2に拾わせたい情報は、リストではなく"
    "ページ本文（テキストWeb パーツ）に記載する必要があります（詳細は「クライアント資料2案比較」シート参照）。",
    height=48, fill="FDEBEC", bold=True,
)
r += 1
r = note_box(
    ws, r, 1, LAST_COL_8,
    "両者は別エージェントとして分離するのが妥当（「文書検索QA」と「ファイル解析＋ルール照合」で用途が全く異なるため、"
    "1つのエージェントに無理に統合すると精度が落ちる）。",
    height=28, italic=True,
)
r += 2

# ---------------------------------------------------------------------------
# 根拠資料一覧（gpt-check#17：参照資料名を実在ファイル名・節番号に統一したうえで一覧化）
# ---------------------------------------------------------------------------
r = section_title(ws, r, 1, LAST_COL_8, "根拠資料一覧（本書が参照する一次情報・実在ファイル名）", size=12)
ref_headers = ["ファイル名", "内容"]
ref_spans = [(1, 5), (6, 14)]
for (a, b), h in zip(ref_spans, ref_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY_80)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 20
r += 1

ref_rows = [
    ("docs/01_current-site-analysis.md", "現状課題22件（T-1〜4/B-1〜7/C-1〜3/CA-1〜2/P-1〜2/G-1〜4）とアナリティクス所見の一次情報。課題・数値の正。"),
    ("差戻し回答書.md", "初回の社長差戻しへの回答（クイックウィン廃止・全課題＋AI-1,2同日公開の方針確定・全ページ横断の課題）。"),
    ("差戻し回答書_v2.md", "2回目の社長差戻しへの回答（4→3カラム、パーツ名修正、アンカーカード整理等）。"),
    ("差戻し回答書_v3.md", "3回目の社長差戻しへの回答（AI-2の外部リンク探索能力・クライアントページのリンク提示可否の確認事項）。"),
    ("差戻し回答書_v4.md", "4回目の社長差戻しへの回答（クライアント資料2案比較、KPIの権限検証、AI-1,2実装方式の確定）。"),
    ("sharepoint-feasibility.md", "SharePoint実現性調査（R1〜R9）。検索ボックス位置(R1)・端末別出し分け(R8)・利用状況分析(R9)等の一次情報。"),
    ("sharepoint-verify-v3.md", "SharePoint UI仕様の精密再検証（V1〜V7）。セクション最大列数・Web パーツ日本語名称・折りたたみ機能等。"),
    ("sharepoint-verify-v4.md", "SharePoint UI仕様の精密再検証（V8〜V13）。クイックリンクのレイアウト・AI-2のリンク探索能力等。"),
    ("sharepoint-verify-v5.md", "SharePoint UI仕様の精密再検証（V14〜V17）。カスタムリストの検索対象可否・利用状況レポートの閲覧権限等。"),
    ("AI-1-2実装調査.md", "AI-1（Copilot Studio）・AI-2（SharePoint agent）の実装方式・確認事項の技術調査。"),
    ("materials/screenshots/*.png", "AsIs注釈画像の元スクリーンショット（PC/SP・5ページ）。"),
    ("wf_v5/build_tobe_wf_v5.py", "ToBeワイヤーフレーム v5（社長承認済み）の生成スクリプト。本書のToBe画像・記述の設計根拠。"),
]
for i, (fname, desc) in enumerate(ref_rows):
    shade = LIGHTBG if i % 2 == 0 else WHITE
    for (a, b), val in zip(ref_spans, (fname, desc)):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        cell.font = F(size=9.5, bold=(a == 1))
        cell.alignment = WRAP_TOP
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=shade)
    ws.row_dimensions[r].height = 32
    r += 1


# ---------------------------------------------------------------------------
# Sheet 9: リスク登録簿
# ---------------------------------------------------------------------------
ws = wb.create_sheet("リスク登録簿")
ws.sheet_view.showGridLines = False
LAST_COL_9 = 17
set_col_widths(ws, 1, LAST_COL_9, width=13)
page_header_bar(ws, 1, 1, LAST_COL_9, "リスク登録簿（意思決定に必要な主要リスク）", 9, TOTAL_SHEETS)
r = note_box(
    ws, 3, 1, LAST_COL_9,
    "根拠: AI-1-2実装調査.md、差戻し回答書_v3.md／_v4.md、sharepoint-verify-v3〜v5.md。"
    "実装可否そのものに関わる／同日カットオーバー公開の前提を左右するリスクを優先して掲載しています。",
    height=32, italic=True,
)
r += 1

risk_headers = ["リスク", "影響", "対応策", "確認先"]
risk_spans = [(1, 4), (5, 8), (9, 13), (14, 17)]
for (a, b), h in zip(risk_spans, risk_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 22
r += 1

risk_rows = [
    ("AI-1のカスタムWebサイトチャネル埋め込みで「認証なし」設定が必要になる可能性",
     "「認証なし」公開が社内ポリシー・情報セキュリティ基準に抵触する場合、AI-1をこの方式で実装できず、代替実装の検討または実装自体の見送りが必要になる。",
     "F0でIT窓口・情報セキュリティに公開設定の許容可否を確認。許容不可の場合はCopilot Studio内の別チャネル（Teams等）や他の埋め込み方式を代替検討する。",
     "SharePoint管理者・IT窓口／情報セキュリティ"),
    ("AI-1のPPTX直接アップロードが実験的機能（Experimental）である",
     "実験的機能は予告なく仕様変更・提供停止される可能性があり、運用が不安定になるリスクがある。",
     "当面はPDF変換を前段に挟む運用を既定とし、実験的機能への依存を避ける。機能の正式提供状況をF2〜F3で再確認する。",
     "実装担当"),
    ("AI-2の外部リンク提示能力が一次情報で未確証",
     "「ページ本文中の外部URLを期待どおり提示できるか」が公式に明記されておらず、クライアント向け外部リンク案内が期待通り機能しない可能性がある。",
     "F3で実機検証を必須タスク化し、期待どおり動作しない場合はページ本文内の記述方法（アンカーテキストの明示化等）を調整する。",
     "実装担当（F3実機検証）"),
    ("M365 Copilot／Copilot Studioのライセンス・予算が確保できない",
     "AI-1,2の実装可否・費用の前提が崩れ、同日カットオーバー公開そのものが成立しなくなる（最重要リスク）。",
     "F0を最優先で実施し、可否・費用を早期確定。確保できない場合は公開延期が原則（Must要件を外した部分公開は行わない）。",
     "社長／SharePoint管理者・IT窓口"),
    ("AI判定（AI-1のブランド準拠チェック）の精度は一次スクリーニング止まりである",
     "色のΔE数値照合やロゴのピクセル単位検出など厳密な自動判定は標準機能で保証されないため、AIの判定のみでブランド逸脱を見逃すリスクがある。",
     "AI-1は「気づきの提示」用途と位置づけ、最終承認は人（ブランドレビュー担当）が行う運用を維持し、品質保証プロセスを残す。",
     "マーケティング責任者"),
]
for i, row_vals in enumerate(risk_rows):
    shade = LIGHTBG if i % 2 == 0 else WHITE
    for (a, b), val in zip(risk_spans, row_vals):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        cell.font = F(size=9.5, bold=(a == 1))
        cell.alignment = WRAP_TOP
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=shade)
    ws.row_dimensions[r].height = 80
    r += 1


# ---------------------------------------------------------------------------
# 印刷設定（PDF書き出し・印刷時に横1ページに収める。Excel通常閲覧には影響しない）
# ---------------------------------------------------------------------------
for _ws in wb.worksheets:
    _ws.page_setup.orientation = "landscape"
    _ws.page_setup.fitToWidth = 1
    _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr = WorksheetProperties(pageSetUpPr=PageSetupProperties(fitToPage=True)).pageSetUpPr
    _ws.page_margins.left = 0.3
    _ws.page_margins.right = 0.3
    _ws.page_margins.top = 0.4
    _ws.page_margins.bottom = 0.4

# ---------------------------------------------------------------------------
# 保存
# ---------------------------------------------------------------------------
os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
wb.save(OUT_PATH)
print("saved", OUT_PATH, round(os.path.getsize(OUT_PATH) / 1024 / 1024, 2), "MB")



