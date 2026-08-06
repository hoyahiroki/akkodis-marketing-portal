#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AKKODiSマーケティングポータル リニューアル提案書（清書版・Excel）生成スクリプト

- 提案として読める文章だけを載せる。制作過程のやり取り・検討経緯は載せない。
- ワイヤーフレームは 1ページ＝1シート。左にWF画像（1枚もの）、右に番号対応のコメント表を置く。
- コメント表の各行は、WF画像内の該当セクションと同じ高さに来るよう行高を計算して配置する。

出力: AKKODiSマーケティングポータル_リニューアル提案書_清書版.xlsx
既存の「AKKODiSマーケティングポータル_リニューアル提案書.xlsx」には一切触れない。
"""
import json
import math
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.properties import PageSetupProperties, WorksheetProperties
from PIL import Image as PILImage

ROOT = "/workspace/akkodis-marketing-portal"
TASK_DIR = os.path.join(ROOT, "deliverables/1784786903168-17720")
ASSET_DIR = os.path.join(TASK_DIR, "xlsx_assets")
CLEAN_ASSET_DIR = os.path.join(ASSET_DIR, "clean")
WF_DIR = os.path.join(TASK_DIR, "wf_v6")
OUT_PATH = os.path.join(TASK_DIR, "AKKODiSマーケティングポータル_リニューアル提案書_清書版.xlsx")

DOC_DATE = "2026-08-06"
DOC_VERSION = "1.0"

# ---------------------------------------------------------------------------
# ブランドトークン（akkodis-brand-core）
# ---------------------------------------------------------------------------
NAVY = "001F33"
YELLOW = "FFB81C"
WHITE = "FFFFFF"
NAVY_80 = "334B5C"
NAVY_40 = "999AAD"
NAVY_20 = "CCC8D6"
LIGHTBG = "F7F7F5"
BLACK = "000000"
ACCENT_RED = "C41E3A"

FONT_NAME = "Meiryo UI"


def F(size=10, bold=False, color=BLACK, italic=False, name=FONT_NAME):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


THIN = Side(style="thin", color=NAVY_40)
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
WRAP_TOP_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

TOTAL_SHEETS = 14


# ---------------------------------------------------------------------------
# 共通パーツ
# ---------------------------------------------------------------------------
def set_col_widths(ws, first_col, last_col, width):
    for c in range(first_col, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = width


def set_rows_height(ws, r0, r1, pt=15):
    for r in range(r0, r1 + 1):
        ws.row_dimensions[r].height = pt


def col_px(width):
    """列幅（文字単位）→ ピクセル（既定フォントの1文字幅7px＋パディング5px）"""
    return round(width * 7) + 5


def page_header_bar(ws, row, first_col, last_col, title, sheet_no):
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col - 2)
    c = ws.cell(row=row, column=first_col,
                value=f"AKKODiSマーケティングポータル リニューアル提案　|　{title}")
    c.font = F(size=15, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(first_col, last_col - 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=row, start_column=last_col - 1, end_row=row, end_column=last_col)
    c2 = ws.cell(row=row, column=last_col - 1, value=f"社外秘　{sheet_no}/{TOTAL_SHEETS}")
    c2.font = F(size=10, bold=True, color=YELLOW)
    c2.alignment = Alignment(horizontal="right", vertical="center", indent=1)
    c2.fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[row].height = 30
    return row + 1


def section_title(ws, row, first_col, last_col, text, fill=NAVY, color=WHITE, size=12, height=24):
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=first_col, value=text)
    c.font = F(size=size, bold=True, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for col in range(first_col, last_col + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)
    ws.row_dimensions[row].height = height
    return row + 1


def note_box(ws, row, first_col, last_col, text, height=40, fill=LIGHTBG, bold=False,
             color=BLACK, size=10, italic=False):
    ws.merge_cells(start_row=row, start_column=first_col, end_row=row, end_column=last_col)
    c = ws.cell(row=row, column=first_col, value=text)
    c.font = F(size=size, bold=bold, color=color, italic=italic)
    c.alignment = WRAP_TOP
    for col in range(first_col, last_col + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill)
        ws.cell(row=row, column=col).border = BORDER_ALL
    ws.row_dimensions[row].height = height
    return row + 1


def table(ws, row, spans, headers, rows, heights, header_fill=NAVY, header_size=10,
          center_cols=(), status_col=None, zebra=True, header_h=22, body_size=10,
          bold_cols=()):
    """汎用テーブル。spans=[(colA,colB),…]、rows=[(値,…),…]、heights=int or list"""
    for (a, b), h in zip(spans, headers):
        ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
        cell = ws.cell(row=row, column=a, value=h)
        cell.font = F(size=header_size, bold=True, color=WHITE)
        cell.alignment = CENTER
        for col in range(a, b + 1):
            ws.cell(row=row, column=col).border = BORDER_ALL
            ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=header_fill)
    ws.row_dimensions[row].height = header_h
    row += 1

    for i, vals in enumerate(rows):
        shade = (LIGHTBG if i % 2 == 0 else WHITE) if zebra else WHITE
        for (a, b), val in zip(spans, vals):
            ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
            cell = ws.cell(row=row, column=a, value=val)
            is_status = status_col is not None and a == status_col
            fill_color = shade
            font_color = BLACK
            if is_status:
                fill_color, font_color = STATUS_COLOR.get(val, (shade, BLACK))
            cell.font = F(size=body_size, bold=(a in bold_cols or is_status), color=font_color)
            cell.alignment = WRAP_TOP_CENTER if (a in center_cols or is_status) else WRAP_TOP
            for col in range(a, b + 1):
                ws.cell(row=row, column=col).border = BORDER_ALL
                ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=fill_color)
        ws.row_dimensions[row].height = heights[i] if isinstance(heights, list) else heights
        row += 1
    return row


# ---------------------------------------------------------------------------
# 区分（5分類）
# ---------------------------------------------------------------------------
S_RESOLVED = "解消"
S_KEEP = "現状維持"
S_OUT = "対象外"
S_GOOD = "良好・維持"
S_CHECK = "要現物確認"

STATUS_COLOR = {
    S_RESOLVED: ("E7F2EA", NAVY_80),
    S_KEEP: ("EFEFEF", NAVY_80),
    S_OUT: ("EFEFEF", NAVY_40),
    S_GOOD: ("FFF4E0", NAVY_80),
    S_CHECK: ("FDEBEC", ACCENT_RED),
}


# ---------------------------------------------------------------------------
# 画像アセットの準備（wf_v6 の清書済みWFから、Excel掲載用の縮小版を作る）
# ---------------------------------------------------------------------------
WF_PAGES = [
    ("top", "1. トップ（ホーム）"),
    ("brand", "2. ブランドライブラリ"),
    ("client", "3. クライアント向け"),
    ("candidate", "4. キャンディデート向け"),
    ("partner", "5. パートナー向け"),
]

WF_DISPLAY_W = 900       # Excel上での表示幅(px)。5枚とも共通
WF_EMBED_W = 1800        # 埋め込み実体の幅(px)。表示幅の2倍で解像度を確保
THUMB_W = 230


def prepare_assets():
    os.makedirs(CLEAN_ASSET_DIR, exist_ok=True)
    made = {}
    for key, _ in WF_PAGES:
        src = os.path.join(WF_DIR, f"{key}.png")
        im = PILImage.open(src)
        big = os.path.join(CLEAN_ASSET_DIR, f"wf_{key}.png")
        im.resize((WF_EMBED_W, round(im.height * WF_EMBED_W / im.width)),
                  PILImage.LANCZOS).save(big, optimize=True)
        th = os.path.join(CLEAN_ASSET_DIR, f"thumb_{key}.png")
        im.resize((THUMB_W, round(im.height * THUMB_W / im.width)),
                  PILImage.LANCZOS).save(th, optimize=True)
        made[key] = (big, th)
    return made


ASSETS = prepare_assets()
COMMENTS = json.load(open(os.path.join(WF_DIR, "comments.json"), encoding="utf-8"))


# ===========================================================================
# Workbook
# ===========================================================================
wb = Workbook()
wb.remove(wb.active)


# ---------------------------------------------------------------------------
# シート1: 表紙
# ---------------------------------------------------------------------------
ws = wb.create_sheet("表紙")
ws.sheet_view.showGridLines = False
set_col_widths(ws, 1, 14, width=10)

logo_path = os.path.join(ASSET_DIR, "AKKODIS_Logo_POS_RGB.png")
lw, lh = PILImage.open(logo_path).size
_img = XLImage(logo_path)
_img.width = 260
_img.height = round(lh * 260 / lw)
ws.add_image(_img, "B3")
set_rows_height(ws, 1, 12, pt=22)

ws.merge_cells("B12:M12")
c = ws.cell(row=12, column=2, value="社外秘（Confidential）")
c.font = F(size=11, bold=True, color=WHITE)
c.fill = PatternFill("solid", fgColor=NAVY)
c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[12].height = 26

ws.merge_cells("B15:N17")
c = ws.cell(row=15, column=2, value="AKKODiSマーケティングポータル\nリニューアル提案")
c.font = F(size=28, bold=True, color=NAVY)
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
set_rows_height(ws, 15, 17, pt=26)

ws.merge_cells("B19:N21")
c = ws.cell(
    row=19, column=2,
    value="現状課題22件の整理と、SharePoint標準機能によるToBe設計（ワイヤーフレーム5ページ）\n"
          "AI-1（ブランドセルフチェッカー）／AI-2（マーケ特化AIアシスタント）の実装をマスト要件として統合",
)
c.font = F(size=13, color=NAVY_80)
c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
set_rows_height(ws, 19, 21, pt=22)

info_rows = [
    ("対象", "AKKODiS マーケティングポータル（SharePoint Online イントラサイト）"),
    ("範囲", "現状課題の整理／ToBeワイヤーフレーム（5ページ）／フェーズ計画／KPI／AI実装方式／リスク"),
    ("作成日", DOC_DATE),
    ("版数", f"v{DOC_VERSION}"),
]
r = 24
for label, val in info_rows:
    ws.cell(row=r, column=2, value=label).font = F(size=10, bold=True, color=WHITE)
    ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=NAVY_80)
    ws.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=13)
    vcell = ws.cell(row=r, column=3, value=val)
    vcell.font = F(size=10)
    vcell.alignment = WRAP_TOP
    for col in range(2, 14):
        ws.cell(row=r, column=col).border = BORDER_ALL
    ws.row_dimensions[r].height = 22
    r += 1

ws.merge_cells(f"B{r + 2}:N{r + 5}")
note = ws.cell(
    row=r + 2, column=2,
    value="本書は社外秘情報を含みます。外部への持ち出し・転載はご遠慮ください。\n"
          "文言が未確定の箇所は「マーケ部確定待ち」と明記しています。"
          "断定できない技術事項は「実機検証で確認する」と明記し、確定事項と区別しています。",
)
note.font = F(size=9.5, italic=True, color=NAVY_80)
note.alignment = WRAP_TOP
set_rows_height(ws, r + 2, r + 5, pt=20)


# ---------------------------------------------------------------------------
# シート2: エグゼクティブサマリー
# ---------------------------------------------------------------------------
ws = wb.create_sheet("エグゼクティブサマリー")
ws.sheet_view.showGridLines = False
LC = 16
set_col_widths(ws, 1, LC, width=12)
r = page_header_bar(ws, 1, 1, LC, "エグゼクティブサマリー", 2)
r += 1

r = section_title(ws, r, 1, LC, "1. 提案の狙い")
for p in [
    "「探している資料が見つからない」という利用実態を解消し、ポータルを起点に自力で目的の情報へ到達できる状態をつくる。",
    "5ページで異なっていた構成の作法を、共通テンプレート（冒頭に目次／見出しはH2で統一／注意喚起はセクション背景色）に揃え、ページを移動するたびに探し方を学び直す状態をなくす。",
    "AI-1（ブランドセルフチェッカー）とAI-2（マーケ特化AIアシスタント）を組み込み、依頼前の自己点検と、自然言語での資料案内を成立させる。",
    "SharePointの標準Web パーツで実現できる範囲を中心に設計し、追加開発を伴わない構成で運用を継続できるようにする。",
]:
    r = note_box(ws, r, 1, LC, "・" + p, height=32)

r += 1
r = section_title(ws, r, 1, LC, "2. 現状課題22件の内訳（ページ別）")
r = table(
    ws, r,
    spans=[(1, 2), (3, 7), (8, 9), (10, 16)],
    headers=["区分", "ページ／範囲", "件数", "課題ID"],
    rows=[
        ("T", "トップ（ホーム）", "4", "T-1〜T-4"),
        ("B", "ブランドライブラリ", "7", "B-1〜B-7"),
        ("C", "クライアント向け", "3", "C-1〜C-3"),
        ("CA", "キャンディデート向け", "2", "CA-1〜CA-2"),
        ("P", "グローバルブランドパートナーシップ", "2", "P-1〜P-2"),
        ("G", "全ページ横断", "4", "G-1〜G-4"),
        ("合計", "", "22", ""),
    ],
    heights=20, center_cols=(1, 3, 8, 10), header_h=20, bold_cols=(1,),
)

r += 1
r = section_title(ws, r, 1, LC, "3. 改善対象の内訳（区分別）", fill=NAVY_80, size=11)
r = table(
    ws, r,
    spans=[(1, 4), (5, 6), (7, 16)],
    headers=["区分", "件数", "内容"],
    rows=[
        (S_RESOLVED, "14", "設計変更により改善する（T-3・T-4／B-1〜B-4・B-6・B-7／C-1・C-2／P-2／G-1〜G-3）"),
        (S_KEEP, "2", "設計を変更せず現状を維持する（T-2＝組織体制の配置／CA-1＝各チーム連絡先の現行踏襲）"),
        (S_OUT, "2", "本構築のスコープ対象外（T-1＝モバイル対応／G-4＝権限により非表示のため対応不要）"),
        (S_GOOD, "2", "既に良好なため他ページへ横展開する（C-3・CA-2）"),
        (S_CHECK, "2", "構成は変更せず実装フェーズで現物を確認する（B-5・P-1）"),
        ("合計", "22", ""),
    ],
    heights=26, center_cols=(1, 5), status_col=1, header_h=20, zebra=False,
)

r += 1
r = section_title(ws, r, 1, LC, "4. 優先着手すべき課題トップ3")
for cid, title, desc in [
    ("B-1", "ブランドライブラリの情報過多・ページ内ナビ欠如",
     "インパクト高／コスト中〜高。直近30日で最も使われる実働ページ（一意閲覧者319名）であり、「ガイドラインが増えてわかりにくい」という利用者の声に直接対応する。改善効果が最も広く届く。"),
    ("G-2", "ポータル内の検索性の弱さ",
     "インパクト高／コスト中〜高。「見つからないのでメールを探す」という状態の根本原因であり、AI-2導入の主戦場となる。"),
    ("C-1", "クライアント向け「お役立ち資料」の30本超のリンク羅列",
     "インパクト高／コスト低〜中。説明文・分類・絞り込みがなく発見コストが高い。資料の並べ方をAI-2が回答根拠にできる形へ変えることで、検索性の改善と両立できる。"),
]:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c1 = ws.cell(row=r, column=1, value=cid)
    c1.font = F(bold=True, color=WHITE)
    c1.fill = PatternFill("solid", fgColor=NAVY_80)
    c1.alignment = CENTER
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=16)
    c2 = ws.cell(row=r, column=3, value=f"{title}\n{desc}")
    c2.font = F(size=10)
    c2.alignment = WRAP_TOP
    for col in range(1, 17):
        ws.cell(row=r, column=col).border = BORDER_ALL
    ws.row_dimensions[r].height = 46
    r += 1

r += 1
r = section_title(ws, r, 1, LC, "5. 改善方針")
for p in [
    "改善対象14課題を SharePoint 標準Web パーツ（クイックリンク／テキスト／行動喚起／埋め込み 等）で実装し、追加ライセンスなしで実現できる範囲を中心に構築する。現状維持・対象外・良好・維持・要現物確認の8課題は設計を変更せず、区分として状態を明示する。",
    "AI-1（ブランドセルフチェッカー：Copilot Studio）／AI-2（マーケ特化AIアシスタント：SharePoint agent）の実装をマスト要件とし、一部機能だけを先行公開する進め方は採らない。",
    "改善対象14課題とAI-1・AI-2を1つのリリースとしてまとめ、同日カットオーバーで公開する（フェーズ F0〜F4。詳細は「フェーズ計画」シート）。",
    "AI-1・AI-2の前提となる M365 Copilot／Copilot Studio のライセンスと管理者許可の確認（F0）が最大のクリティカルパス。F0でライセンス・権限が確保できない場合は同日公開の前提が崩れるため、公開延期を原則とする。",
]:
    r = note_box(ws, r, 1, LC, "・" + p, height=34)

r += 1
r = section_title(ws, r, 1, LC, "6. 全体像（ToBe ワイヤーフレーム 5ページ）")
r += 1
thumb_row = r
thumb_cols = [2, 5, 8, 11, 14]
max_h = 0
for col, (key, _) in zip(thumb_cols, WF_PAGES):
    p = ASSETS[key][1]
    w, h = PILImage.open(p).size
    im = XLImage(p)
    ws.add_image(im, f"{get_column_letter(col)}{thumb_row}")
    max_h = max(max_h, h)
rows_needed = math.ceil(max_h / 20) + 2
set_rows_height(ws, thumb_row, thumb_row + rows_needed)
label_row = thumb_row + rows_needed + 1
for col, (_, label) in zip(thumb_cols, WF_PAGES):
    ws.merge_cells(start_row=label_row, start_column=col, end_row=label_row, end_column=col + 2)
    c = ws.cell(row=label_row, column=col, value=label)
    c.font = F(size=9, bold=True, color=NAVY)
    c.alignment = Alignment(horizontal="center")
r = label_row + 2
r = note_box(ws, r, 1, LC,
             "※各ページの構成と、セクションごとの設計意図は「WF 1.〜WF 5.」の各シートを参照してください。",
             height=24)


# ---------------------------------------------------------------------------
# シート3: 課題一覧（全22課題）
# ---------------------------------------------------------------------------
ISSUES = [
    # (ID, 区分, ページ, 課題（現状）, 改善内容（ToBe）, 期待効果)
    ("T-1", S_OUT, "トップ",
     "埋め込みダッシュボードがモバイル非対応。スマートフォン表示では極小枠と横スクロールになり、指標・グラフが実質判読できない。",
     "本構築のスコープ対象外とし、PC表示を対象とする。端末別のコンテンツ出し分けはSharePointのモダンページ標準機能では実現できないため、ワイヤーフレームにも描かない。",
     "スコープが明確になり、モバイル対応の要否は次フェーズの判断事項として引き継げる（対応漏れではなく意図的な対象外）。"),
    ("T-2", S_KEEP, "トップ",
     "「組織体制」が最下部に配置され、来訪者が最初に求める「探す・依頼する」導線より下位に沈んでいる。",
     "配置は最下部のまま変更しない（現状維持・方針決定済み）。冒頭の目次（クイックリンク Web パーツ）から組織体制へのジャンプ導線のみ追加する。",
     "配置変更のコストをかけずに、目次からワンクリックで到達できるようになる。"),
    ("T-3", S_RESOLVED, "トップ",
     "4カテゴリカードとグローバルナビが重複している。カードに説明文がなく「次に何ができるか」が伝わらない。",
     "各カード直下に1行の説明文欄を設ける。クイックリンク Web パーツで4件を横並びに配置する（列数は画面幅により変動するため、実機プレビューで確認する）。",
     "クリック前に遷移先の内容が把握でき、期待値を形成できる（説明文の文面はマーケ部確定待ち）。"),
    ("T-4", S_RESOLVED, "トップ",
     "「業務依頼フォーム」「ブランドレビューの手順」など複数の入口が並び、使い分けが1画面で判別しづらい。",
     "「①AI-1セルフチェック→②手順ガイド→③業務依頼フォーム」を上から下へ全幅で並べる縦積み3ステップ導線に再設計する。各ページに共通の窓口案内ブロックを置き、入口を集約する（B-3のフォーム→Teams→Salesforce、CA-1の各チーム連絡といった下流プロセスは現行を踏襲する）。",
     "依頼前にAI-1で気づきを得たうえで正しい手順に沿って依頼でき、ブランドレビューでの再提出の削減が期待できる。"),
    ("B-1", S_RESOLVED, "ブランドライブラリ",
     "情報過多・深い階層・ページ内ナビ（目次／アンカー）の欠如。1ページに全ガイドラインと全手続きを積層しており、目的の情報に辿り着けない。",
     "冒頭にクイックリンク Web パーツによる目次を新設する。ガイドラインライブラリはセクションの折りたたみ機能で格納する（既定は折りたたみ）。",
     "直近30日で最も利用されるページの探索コストを削減し、「わかりにくい」という利用者の声に直接応える。"),
    ("B-2", S_RESOLVED, "ブランドライブラリ",
     "ガイドライン約10本がフラットに列挙されている。粒度・目的（ロゴ／写真／動画／文字表記／AI画像 等）で束ねられていない。",
     "目的軸の4カテゴリ（全社ブランド規定／制作物別／表記・用語／Akkodis Intelligence）にテキスト Web パーツで再構成する。新規ライブラリは作らず、既存資料へのリンクを整理する。",
     "目的別に探せるようになり、10本を都度スキャンする負担がなくなる。"),
    ("B-3", S_RESOLVED, "ブランドライブラリ",
     "ロゴ掲載フローがフォーム→Teams→Salesforceをまたぐ6ステップの長文で説明されている。手続き型コンテンツをテキストのみで表現しており実行負荷が高い。",
     "①先方ロゴ掲載／②自社ロゴ提供の2ケースを2カラムセクションで並置し、番号付きリストで手順を整理する。各カラム末尾に同意書テンプレートへの直リンクを置き、AI-1セルフチェック導線を冒頭に組み込む。",
     "手続きの見通しが良くなり実行負荷が下がる。事前のセルフチェックにより申請の再提出が減ることが期待できる。"),
    ("B-4", S_RESOLVED, "ブランドライブラリ",
     "資料の命名規則が不統一（Akkodis／AKKODiSの表記ゆれ、版・言語・形式の混在）で、一覧の走査性を下げている。",
     "表示名・メタデータをAKKODiS表記に統一する運用ルールを列設計として明記する。",
     "一覧の走査性が向上し、表記ぶれによる混乱・誤選択を防止できる。"),
    ("B-5", S_CHECK, "ブランドライブラリ",
     "ページ内ナビゲーション用のカード3枚について、ラベルの有無とアフォーダンスの確認が必要。",
     "ラベル付きのページ内アンカーリンクカードとして機能しているため、構成を変更しない。実装フェーズで現物を確認する。",
     "不要な改修を行わず、現行の機能をそのまま維持できる。"),
    ("B-6", S_RESOLVED, "ブランドライブラリ",
     "冒頭見出しの黄色マーカーが不統一で、帯見出しとの二重表現になっている。体裁が粗く見える。",
     "モダンページのテキストハイライト色は固定パレットからの選択に限られるため、黄色マーカーは用いず、注意喚起はセクション背景色に統一する。",
     "ページ全体の視覚言語が統一され、全ページ共通のトーンに揃う。"),
    ("B-7", S_RESOLVED, "ブランドライブラリ",
     "スマートフォン表示が極端に長尺（1万px超）で、モバイルでの回遊が困難。",
     "B-1〜B-6の再構成（目次の新設／折りたたみ格納／4カテゴリ整理）により、ページの実効的な情報量と全長を圧縮する。抜本的なモバイル最適化は次フェーズで検討する。",
     "情報整理によりスマートフォン表示の実効的な長さも縮小する。抜本対応は次フェーズに引き継ぐため、改善度合いは限定的である。"),
    ("C-1", S_RESOLVED, "クライアント向け",
     "「お役立ち資料」が30本超のリンク羅列になっている。説明文・サムネイル・絞り込みがなく、目的の資料の発見コストが高い。",
     "公開資料ライブラリをテキスト Web パーツでカテゴリ見出し＋リンク箇条書きに再構成し、30本超はセクションの折りたたみで格納する（推奨案）。リストを併用する構成も比較のうえ併記する（詳細は「クライアント公開資料 構成比較」シート）。",
     "推奨案ではAI-2が資料リンクを回答根拠にできる。リストを併用する構成ではカテゴリ・言語での絞り込みや並べ替えが可能になる。"),
    ("C-2", S_RESOLVED, "クライアント向け",
     "リンクに更新日・対象・要約などのメタ情報がなく、最新版か・誰向けかが判別できない。",
     "更新日は「日付を出せない資料がある」事情を尊重して任意項目のまま維持し、テキスト Web パーツ内の見出しとカテゴリ分類で対象を明確化する。",
     "全件必須化を強制せずに運用を維持しながら、カテゴリ単位での対象・種別の視認性を確保する（改善度合いは限定的）。"),
    ("C-3", S_GOOD, "クライアント向け",
     "冒頭のアンカー型クイックリンク3件が機能しており、他ページより回遊への配慮がある。一方でこの良い部分が他ページへ展開されていない。",
     "「冒頭に目次（クイックリンク）を置く」パターンを全ページ共通テンプレートとして横展開する（横断課題G-1として全ページに適用）。",
     "ページ間で「探し方」を学び直す必要がなくなり、サイト全体の回遊性が向上する。"),
    ("CA-1", S_KEEP, "キャンディデート向け",
     "マスター資料が「各チームに直接連絡してください」で完結しており、資料そのものが入手できない。「見つからないのでメールを探す」を誘発する典型例。",
     "現行の文言（各チーム連絡先）をテキスト Web パーツでそのまま維持する（現状維持・方針決定済み）。直リンク化が可能と確認できた場合は再設計の余地を残す。",
     "意図的な現状維持とし、確認結果によっては次段階で見直せる余地を残す。"),
    ("CA-2", S_GOOD, "キャンディデート向け",
     "5ページ中もっとも簡潔で見やすい構成（AKKODiS公式HP／マスター資料／Social Mediaの3セクション）になっている。",
     "他ページ再設計の参照モデルとしてこの3構成を踏襲する。Social Mediaはクイックリンク Web パーツで4件の横並びとし（列数は実機プレビューで確認）、AKKODiS Peopleはテキスト Web パーツで表現する。",
     "改修コストを抑えつつ、他ページとの体験の一貫性を確保する参照モデルとして機能する。"),
    ("P-1", S_CHECK, "パートナーシップ",
     "Akkodis ASP Teamのみ余分な画像枠がプレースホルダーとして表示され、他2チームには存在しない枠が体裁を崩している。",
     "画像枠を是正する（差し替えまたは除去）。他の改善と合わせて同日カットオーバーで実施し、実装フェーズで現物を確認のうえ対応する。",
     "3パートナーの体裁が揃い、見た目の一貫性が回復する。"),
    ("P-2", S_RESOLVED, "パートナーシップ",
     "長尺ページだがページ内ナビがなく、クライアント向けページのアンカー導線と不統一。",
     "ページ内目次（クイックリンク Web パーツ）を新設する。3パートナーのサムネイルカードは目次と役割が重複するため削除し、目次に一本化する。",
     "クライアント向けページと同様の回遊導線を確保し、横断課題G-1の解消に寄与する。"),
    ("G-1", S_RESOLVED, "全ページ横断",
     "ページ間で構成が不統一（アンカーリンクの有無、装飾の使い方、情報密度）。回遊のたびに「探し方」を学び直す必要がある。",
     "全5ページで「目次を冒頭に配置／見出しはH2で統一／注意喚起はセクション背景色に統一」の共通テンプレートを適用する。クライアント向けページのアンカー導線パターンをブランドライブラリ・パートナーシップへ横展開する。",
     "サイト全体で操作の作法が揃い、ページを移動しても同じ探し方が通用する。"),
    ("G-2", S_RESOLVED, "全ページ横断",
     "検索性が弱い。全ページ左上にSharePoint標準検索はあるが、ページ内の絞り込み・ファセット・タグがない。「検索性が悪い／見つからない」という声の核心。",
     "本文への検索ボックス追加や位置変更は仕様上できないため、ヘッダーの既定検索を維持する。そのうえでAI-2（SharePoint agent）を導入し、自然言語での案内と出典リンク提示で発見性を補う（マスト要件）。",
     "「探す」手段が検索窓だけでなくなり、言葉で尋ねて出典付きで案内される経路が加わる。"),
    ("G-3", S_RESOLVED, "全ページ横断",
     "窓口・相談導線が分散している（業務依頼フォーム／ブランドレビュープロセス／各チームへ直接連絡／Salesforceサービスリクエスト／Teams連絡）。",
     "トップの「マーケへの業務依頼」（AI-1セルフチェック→手順ガイド→業務依頼フォームの縦積み3ステップ）を軸に、各ページへ共通の窓口案内ブロックを置き入口を集約する（業務プロセス自体の一本化ではなく、下流プロセスは現行を維持する）。",
     "「どこから頼めばよいか」の迷いが減り、依頼が正しい窓口へ集まる。"),
    ("G-4", S_OUT, "全ページ横断",
     "ナビゲーションに「編集 のナビゲーション」が露出している。編集者ビュー由来と見られる。",
     "権限のないユーザーには表示されないため対応は不要（本構築のスコープ対象外）。F0の運用確認事項として実機で確認する。",
     "不要な改修を行わず、F0の確認項目として取り扱う。"),
]

ws = wb.create_sheet("課題一覧")
ws.sheet_view.showGridLines = False
LC = 20
ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 8
ws.column_dimensions["C"].width = 13
set_col_widths(ws, 4, LC, width=12)
r = page_header_bar(ws, 1, 1, LC, "課題一覧（全22課題・区分付き）", 3)
r += 1
r = note_box(
    ws, r, 1, LC,
    "区分は 解消／現状維持／対象外／良好・維持／要現物確認 の5つで、課題ごとに扱いを明示しています。"
    "「全22課題を解消」という一括表現は用いていません（区分別の件数はエグゼクティブサマリーを参照）。",
    height=30, italic=True,
)
r += 1
r = table(
    ws, r,
    spans=[(1, 1), (2, 2), (3, 3), (4, 8), (9, 14), (15, 20)],
    headers=["課題ID", "区分", "対象ページ", "課題（現状）", "改善内容（ToBe）", "期待効果"],
    rows=ISSUES,
    heights=[96] * len(ISSUES),
    center_cols=(1, 3), status_col=2, bold_cols=(1,), body_size=9.5,
)


# ---------------------------------------------------------------------------
# シート4〜8: WF（1ページ＝1シート・左WF／右コメント）
# ---------------------------------------------------------------------------
WF_PURPOSE = {
    "top": "ポータルの入口として「探す」「依頼する」の2つの導線を1画面で示すページ。"
           "冒頭の目次で全体像を提示し、マーケへの業務依頼は縦積みの3ステップに集約します。",
    "brand": "利用が最も多いページとして、情報量を保ったまま探索コストを下げるページ。"
             "冒頭の目次、目的軸のカテゴリ整理、折りたたみによる格納で、目的の資料へ最短で到達できるようにします。",
    "client": "クライアント向けの公開資料を、AI-2が回答根拠として使える形で整理するページ。"
              "カテゴリ見出しとリンク箇条書きに再構成し、件数の多いブロックは折りたたみで格納します。",
    "candidate": "5ページ中もっとも簡潔な現行構成を維持しつつ、他ページへ展開する参照モデルとして整えるページ。"
                 "3セクション構成を崩さず、パーツの選び方だけを共通テンプレートに揃えます。",
    "partner": "長尺ページの回遊性をページ内目次で確保し、3パートナーの表示体裁を揃えるページ。"
               "目次と役割が重複するサムネイルカードは置きません。",
}

# 列設計（px は既定フォント換算）
WF_A_W = 1.5                      # 左余白
WF_IMG_COLS = (2, 7)              # B〜G: WF画像を貼る領域
WF_IMG_COL_W = 21.9               # 1列あたり 158px × 6列 = 948px（画像900px＋余白）
WF_NO_W, WF_SEC_W, WF_PART_W, WF_ISS_W = 4.5, 20.0, 22.0, 8.0
WF_CMT_COLS = (12, 16)            # L〜P: コメント
WF_CMT_COL_W = 28.0               # 201px × 5列 = 1005px
WF_LAST_COL = 16

MAX_ROW_PT = 405.0                # Excelの行高上限 409.5pt に対する安全値
WF_SCALE = None                   # 実行時に決定（900 / 元画像幅）

_cmt_px = (WF_CMT_COLS[1] - WF_CMT_COLS[0] + 1) * col_px(WF_CMT_COL_W)
_sec_px = col_px(WF_SEC_W)
_part_px = col_px(WF_PART_W)


def _text_px(text, box_px, font_pt):
    """全角基準でおおまかな折り返し行数を見積もり、必要な行高(px)を返す。"""
    char_px = font_pt * 4 / 3
    per_line = max(1, int(box_px / char_px))
    lines = 0
    for para in str(text).split("\n"):
        lines += max(1, math.ceil(len(para) / per_line))
    return lines * char_px * 1.5 + 10


ALIGN_REPORT = []


def build_wf_sheet(key, sheet_no, sheet_name, page_label):
    global WF_SCALE
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = WF_A_W
    set_col_widths(ws, WF_IMG_COLS[0], WF_IMG_COLS[1], width=WF_IMG_COL_W)
    ws.column_dimensions[get_column_letter(8)].width = WF_NO_W
    ws.column_dimensions[get_column_letter(9)].width = WF_SEC_W
    ws.column_dimensions[get_column_letter(10)].width = WF_PART_W
    ws.column_dimensions[get_column_letter(11)].width = WF_ISS_W
    set_col_widths(ws, WF_CMT_COLS[0], WF_CMT_COLS[1], width=WF_CMT_COL_W)

    r = page_header_bar(ws, 1, 1, WF_LAST_COL, f"ToBe ワイヤーフレーム　{page_label}", sheet_no)
    r += 1
    r = note_box(ws, r, 1, WF_LAST_COL, WF_PURPOSE[key], height=34, fill="FFF4E0")
    r = note_box(ws, r, 1, WF_LAST_COL,
                 "左がワイヤーフレーム、右が各セクションへのコメントです。画像内の丸番号と右表の「No」が対応しています。"
                 "ワイヤーフレームは構成を示すもので、文言・画像は実装時に確定します。",
                 height=26, italic=True)
    r += 1

    # 見出し行（左＝WF、右＝コメント表のヘッダー）
    head_row = r
    ws.merge_cells(start_row=head_row, start_column=WF_IMG_COLS[0],
                   end_row=head_row, end_column=WF_IMG_COLS[1])
    c = ws.cell(row=head_row, column=WF_IMG_COLS[0], value="ワイヤーフレーム")
    c.font = F(size=11, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(WF_IMG_COLS[0], WF_IMG_COLS[1] + 1):
        ws.cell(row=head_row, column=col).fill = PatternFill("solid", fgColor=NAVY_80)

    cmt_spans = [(8, 8), (9, 9), (10, 10), (11, 11), WF_CMT_COLS]
    for (a, b), h in zip(cmt_spans, ["No", "セクション", "Web パーツ", "該当課題", "コメント"]):
        ws.merge_cells(start_row=head_row, start_column=a, end_row=head_row, end_column=b)
        cell = ws.cell(row=head_row, column=a, value=h)
        cell.font = F(size=10, bold=True, color=WHITE)
        cell.alignment = CENTER
        for col in range(a, b + 1):
            ws.cell(row=head_row, column=col).border = BORDER_ALL
            ws.cell(row=head_row, column=col).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[head_row].height = 22
    r += 1

    # --- WF画像 ---------------------------------------------------------
    img_row = r
    img_path = ASSETS[key][0]
    iw, ih = PILImage.open(img_path).size
    disp_w = WF_DISPLAY_W
    disp_h = round(ih * disp_w / iw)
    im = XLImage(img_path)
    im.width, im.height = disp_w, disp_h
    ws.add_image(im, f"{get_column_letter(WF_IMG_COLS[0])}{img_row}")

    data = COMMENTS[key]
    src_w = data.get("image_width") or 2320
    scale = disp_w / src_w
    WF_SCALE = scale
    secs = data["sections"]

    cur = img_row
    consumed_px = 0.0     # img_row からの実際の累積ピクセル
    placed = []

    def emit_filler(px):
        nonlocal cur, consumed_px
        while px > 0.4:
            h_px = min(px, MAX_ROW_PT / 0.75)
            pt = round(h_px * 0.75, 1)
            ws.row_dimensions[cur].height = pt
            consumed_px += pt / 0.75
            cur += 1
            px -= h_px

    # 先頭セクションまでの余白（タイトルバー部分）
    emit_filler(secs[0]["y"] * scale)

    for i, s in enumerate(secs):
        target_px = s["y"] * scale                       # 画像上端からの理想位置
        nxt_y = secs[i + 1]["y"] if i + 1 < len(secs) else data["image_height"]
        gap_px = (nxt_y - s["y"]) * scale

        need = max(
            _text_px(s["comment"], _cmt_px, 9),
            _text_px(s["title"], _sec_px, 9),
            _text_px(s["part"], _part_px, 9),
            34,
        )
        content_px = min(need, gap_px) if need <= gap_px else need
        content_px = min(content_px, MAX_ROW_PT / 0.75)

        row = cur
        vals = [str(s["no"]), s["title"], s["part"], s["issue"] or "—", s["comment"]]
        shade = LIGHTBG if i % 2 == 0 else WHITE
        for (a, b), val in zip(cmt_spans, vals):
            ws.merge_cells(start_row=row, start_column=a, end_row=row, end_column=b)
            cell = ws.cell(row=row, column=a, value=val)
            if a == 8:
                cell.font = F(size=11, bold=True, color=WHITE)
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif a == 9:
                cell.font = F(size=9.5, bold=True, color=NAVY)
                cell.alignment = WRAP_TOP
            elif a == 11:
                cell.font = F(size=9, bold=True, color=NAVY_80)
                cell.alignment = WRAP_TOP_CENTER
            else:
                cell.font = F(size=9)
                cell.alignment = WRAP_TOP
            for col in range(a, b + 1):
                ws.cell(row=row, column=col).border = BORDER_ALL
                ws.cell(row=row, column=col).fill = PatternFill(
                    "solid", fgColor=(NAVY_80 if a == 8 else shade))

        pt = round(content_px * 0.75, 1)
        ws.row_dimensions[row].height = pt
        actual_px = consumed_px
        consumed_px += pt / 0.75
        cur += 1
        placed.append((s["no"], target_px, actual_px, row))

        emit_filler(max(0.0, gap_px - pt / 0.75))

    # 画像下端に合わせて表の下に注記を置く
    r = cur + 1
    r = note_box(ws, r, 1, WF_LAST_COL,
                 "※「該当課題」欄の「—」は、課題一覧に紐づく課題がなく、全ページ共通の設計として置くセクションであることを示します。",
                 height=24, italic=True)

    drifts = [abs(a - t) for _, t, a, _ in placed]
    ALIGN_REPORT.append({
        "sheet": sheet_name,
        "sections": len(secs),
        "scale": round(scale, 5),
        "disp_h": disp_h,
        "max_drift_px": round(max(drifts), 2),
        "rows_used": cur - img_row,
        "detail": [(n, round(t, 1), round(a, 1), round(a - t, 2)) for n, t, a, _ in placed],
    })
    return ws


for i, (key, label) in enumerate(WF_PAGES):
    build_wf_sheet(key, 4 + i, f"WF {label}", label)


# ---------------------------------------------------------------------------
# シート9: クライアント公開資料 構成比較
# ---------------------------------------------------------------------------
ws = wb.create_sheet("クライアント公開資料 構成比較")
ws.sheet_view.showGridLines = False
LC = 14
set_col_widths(ws, 1, LC, width=13)
r = page_header_bar(ws, 1, 1, LC, "クライアント公開資料：3構成の比較", 9)
r += 1
r = note_box(
    ws, r, 1, LC,
    "クライアント向けページの公開資料を、①テキストのみ／②テキスト＋リスト併用／③リストのみ の3構成で比較します。"
    "AI-1・AI-2の実装がマスト要件であるため、「AI-2が資料リンクを回答根拠にできるか」を最重要の判断軸としています。",
    height=34, italic=True,
)
r += 1

r = section_title(ws, r, 1, LC, "結論", size=12)
r = note_box(ws, r, 1, LC,
             "推奨案は ①テキストのみ です（AI-2の回答根拠として使えることを優先）。"
             "②テキスト＋リスト併用 と ③リストのみ は、カテゴリ・言語での絞り込みが必要な場合に要件に応じて選択いただく構成です。"
             "ただし③リストのみはAI-2の回答根拠に使えないため、AI-2の活用を前提とする限り推奨しません。",
             height=40, fill="FFF4E0", bold=True)
r += 1

r = section_title(ws, r, 1, LC, "観点別の比較")
compare_spans = [(1, 2), (3, 5), (6, 6), (7, 9), (10, 10), (11, 13), (14, 14)]
compare_headers = ["観点", "① テキストのみ（推奨）", "判定", "② テキスト＋リスト併用", "判定", "③ リストのみ", "判定"]
for (a, b), h in zip(compare_spans, compare_headers):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    cell = ws.cell(row=r, column=a, value=h)
    cell.font = F(size=9, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = CENTER
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
ws.row_dimensions[r].height = 30
r += 1

judge_cols = (6, 10, 14)
compare_rows = [
    ("① AI-2（生成AI）が回答根拠にできるか",
     "○　ページ本文としてSharePoint agentのナレッジソースに含まれ、回答の根拠になる。", "○",
     "○　同じリンクをテキスト側にも残すため、AI-2はテキスト側を参照して回答根拠にできる。", "○",
     "×　SharePoint agentはリストのデータを使わないとMicrosoft公式ドキュメントに明記されており、回答に出てこない。", "×"),
    ("② カテゴリ／言語の絞り込み・並べ替え",
     "×　手作業の見出し分けのみで、動的な絞り込みはできない。", "×",
     "○　リスト側の列・ビューで絞り込み・並べ替えができる。", "○",
     "○　列・ビューによる絞り込み・並べ替えができる（既定機能）。", "○"),
    ("③ 件数増加時の管理のしやすさ／二重管理コスト",
     "△　本文が長くなり編集が煩雑になりやすい。", "△",
     "×　同じリンクをテキストとリストの両方に維持する必要があり、二重管理コストと更新不整合のリスクが生じる。", "×",
     "○　列で構造的に管理でき、件数増加に強い。", "○"),
    ("④ 人によるサイト内検索でヒットするか",
     "○　ページ本文としてMicrosoft Searchにインデックスされる。", "○",
     "○　テキスト・リストの双方がインデックス対象。", "○",
     "○　リストアイテムも既定でインデックス対象（詳細設定で除外していないことが前提）。", "○"),
    ("⑤ 実装／運用コスト",
     "○　テキスト Web パーツの編集のみで完結し、新規リスト作成が不要で低コスト。", "○",
     "△　新規リスト作成に加えて同一情報をテキスト側にも重複記載・維持する手間が増え、3構成中もっとも運用コストが高い。", "△",
     "△　新規リストの作成・列／ビュー設計・カテゴリ運用ルールの整備など運用の手間が増える。", "△"),
]
for vals in compare_rows:
    for (a, b), val in zip(compare_spans, vals):
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        cell = ws.cell(row=r, column=a, value=val)
        is_judge = a in judge_cols
        cell.font = F(bold=is_judge, size=13 if is_judge else 9,
                      color=(NAVY_80 if val == "○" else (ACCENT_RED if val == "×" else BLACK))
                      if is_judge else BLACK)
        cell.alignment = CENTER if is_judge else WRAP_TOP
        for col in range(a, b + 1):
            ws.cell(row=r, column=col).border = BORDER_ALL
            ws.cell(row=r, column=col).fill = PatternFill(
                "solid", fgColor=LIGHTBG if a == 1 else WHITE)
    ws.row_dimensions[r].height = 58
    r += 1

r += 1
r = note_box(
    ws, r, 1, LC,
    "②案の二重管理コストについて：②テキスト＋リスト併用は、同じ資料リンクをテキスト Web パーツとリストの両方に登録・維持する必要があります。"
    "資料の追加・更新・削除のたびに2箇所を同時に直す運用が求められ、片方の更新を忘れるとテキスト側とリスト側で内容が食い違います。"
    "絞り込み機能とのトレードオフとして、この運用負荷を許容できるかが③との分岐点です。",
    height=46, fill="FDEBEC", bold=True,
)

r += 1
r = section_title(ws, r, 1, LC, "メリット／デメリット")
sum_spans = [(1, 4), (5, 9), (10, 14)]
for (a, b), h in zip(sum_spans, ["① テキストのみ（推奨）", "② テキスト＋リスト併用", "③ リストのみ"]):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    c = ws.cell(row=r, column=a, value=h)
    c.font = F(bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY_80)
    c.alignment = CENTER
ws.row_dimensions[r].height = 20
r += 1
sum_vals = [
    "メリット：AI-2が資料リンクを回答根拠にできる（マスト要件のAI活用と両立する）。実装・運用コストが最も低い。\n"
    "デメリット：カテゴリ・言語での動的な絞り込み・並べ替えができない。",
    "メリット：AI-2の回答根拠を維持しつつ、リスト側で絞り込み・並べ替えができる。\n"
    "デメリット：テキストとリストの二重管理コストと更新不整合のリスクを常時抱える。3構成中もっとも運用コストが高い。",
    "メリット：カテゴリ・言語での絞り込み・並べ替えができ、件数増加時も構造的に管理しやすい。\n"
    "デメリット：AI-2の回答根拠には使われない（Microsoft公式ドキュメントに明記された制約）。新規リスト作成・運用の手間が増える。",
]
for (a, b), val, fill in zip(sum_spans, sum_vals, ["E7F2EA", "FFF4E0", "FDEBEC"]):
    ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
    c = ws.cell(row=r, column=a, value=val)
    c.font = F(size=9.5)
    c.alignment = WRAP_TOP
    for col in range(a, b + 1):
        ws.cell(row=r, column=col).border = BORDER_ALL
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill)
ws.row_dimensions[r].height = 84
r += 2

r = note_box(
    ws, r, 1, LC,
    "将来、リストがSharePoint agentのナレッジソースとして正式に対応された場合は、この比較を見直す価値があります。"
    "本書時点のMicrosoft公式ドキュメントでは未対応です。",
    height=30, italic=True,
)


# ---------------------------------------------------------------------------
# シート10: 横断課題
# ---------------------------------------------------------------------------
ws = wb.create_sheet("横断課題")
ws.sheet_view.showGridLines = False
LC = 14
set_col_widths(ws, 1, LC, width=13)
r = page_header_bar(ws, 1, 1, LC, "全ページ横断の課題（G-1〜G-4）", 10)
r += 1
r = note_box(ws, r, 1, LC,
             "個別ページではなく、5ページ全体に共通して現れる課題と、その対応方針をまとめています。",
             height=24, italic=True)
r += 1
r = table(
    ws, r,
    spans=[(1, 2), (3, 4), (5, 7), (8, 9), (10, 14)],
    headers=["課題ID", "区分", "課題", "観察された状態", "対応方針"],
    rows=[
        ("G-1", S_RESOLVED,
         "ページ間で構成が不統一（アンカーリンクの有無、装飾の使い方、情報密度）。回遊のたびに「探し方」を学び直す必要がある。",
         "クライアント向けはアンカーあり、ブランドライブラリ・パートナーシップはなし。ブランドライブラリの見出しにマーカー装飾が残る。",
         "全5ページで「目次を冒頭に配置／見出しはH2で統一／注意喚起はセクション背景色に統一」の共通テンプレートを適用する。クライアント向けページのアンカー導線パターンをブランドライブラリ・パートナーシップへ横展開する。"),
        ("G-2", S_RESOLVED,
         "検索性が弱い。全ページ左上にSharePoint標準検索はあるが、ページ内の絞り込み・ファセット・タグがない。「検索性が悪い／見つからない」という声の核心。",
         "全ページ共通ヘッダーに標準検索は存在するが、ページ内の絞り込み・ファセットUIはない。",
         "本文への検索ボックス追加や位置変更はSharePointの仕様上できないため、ヘッダーの既定検索を維持する。そのうえでAI-2（SharePoint agent）を導入し、自然言語での案内と出典リンク提示で発見性を補う（マスト要件）。"),
        ("G-3", S_RESOLVED,
         "窓口・相談導線が分散している（業務依頼フォーム／ブランドレビュープロセス／各チームへ直接連絡／Salesforceサービスリクエスト／Teams連絡）。",
         "ブランドライブラリ・キャンディデート向け・トップの3ページに、それぞれ別の窓口が置かれている。",
         "トップの「マーケへの業務依頼」（AI-1セルフチェック→手順ガイド→業務依頼フォームの縦積み3ステップ）を軸に、各ページへ共通の窓口案内ブロックを置き入口を集約する。業務プロセス自体の一本化ではなく、下流プロセス（フォーム→Teams→Salesforce、各チーム連絡）は現行を維持する。"),
        ("G-4", S_OUT,
         "ナビゲーションに「編集 のナビゲーション」が露出している。編集者ビュー由来と見られる。",
         "全ページのナビ末尾に表示。",
         "権限のないユーザーには表示されないため対応は不要（本構築のスコープ対象外）。F0の運用確認事項として実機で確認する。"),
    ],
    heights=[100, 100, 100, 76],
    center_cols=(1,), status_col=3, bold_cols=(1,), body_size=9.5,
)


# ---------------------------------------------------------------------------
# シート11: フェーズ計画
# ---------------------------------------------------------------------------
ws = wb.create_sheet("フェーズ計画")
ws.sheet_view.showGridLines = False
LC = 19
set_col_widths(ws, 1, LC, width=13)
r = page_header_bar(ws, 1, 1, LC, "フェーズ計画（改善対象14課題＋AI-1・AI-2 同日カットオーバー公開）", 11)
r += 1
r = note_box(
    ws, r, 1, LC,
    "前提：改善対象14課題とAI-1・AI-2の実装はマスト要件です（現状維持・対象外・良好・維持・要現物確認の8課題は設計を変更しません）。"
    "一部機能だけを先行して公開する進め方は採らず、1つのリリースとして同日カットオーバーで公開します。"
    "期間はいずれも目安であり、実際の所要はF0の確認結果と体制により変動します。",
    height=40, italic=True,
)
r += 1
r = note_box(
    ws, r, 1, LC,
    "クリティカルパス：F0のAIライセンス確定。AI-1・AI-2は Copilot Studio／M365 Copilot 等のライセンスと管理者許可に依存します。"
    "F0でライセンス・権限が確保できない場合は同日公開の前提が崩れるため、公開延期を原則とします（Must要件を外した部分公開は行いません）。"
    "ライセンス単価はF0で確定・要見積とし、空欄にせず「F0で確定」と明記します。IT窓口の特定が最優先です。",
    height=46, fill="FFF4E0", bold=True,
)
r += 1
r = table(
    ws, r,
    spans=[(1, 2), (3, 4), (5, 8), (9, 10), (11, 13), (14, 17), (18, 19)],
    headers=["フェーズ", "期間目安", "主な作業", "成果物", "完了条件", "承認ゲート（Go／No-Go）", "主担当"],
    rows=[
        ("F0　テナント・ライセンス棚卸し\n（クリティカルパス）",
         "目安2〜4週間\n（ライセンス確認の所要により変動）",
         "M365 Copilot／Copilot Studio の契約・ライセンス割当の確認、DLP／機密ラベル設定の確認、SharePoint管理者・IT窓口の特定、ライセンス単価の見積取得。",
         "ライセンス・機能可否一覧、AI-1・AI-2の実装可否表、ライセンス単価見積",
         "AI-1・AI-2の実装可否と費用（ライセンス単価）が確定する",
         "Go／No-Go①：AI-1・AI-2のライセンス・権限が確保できるか。確保できない場合は同日公開の前提が崩れるため公開延期を原則とする。",
         "マーケティング責任者／SharePoint管理者・IT窓口"),
        ("F1　要件定義・情報設計",
         "目安2〜3週間",
         "全ページのToBeとAI-1・AI-2の要件・権限（所有者／管理者の切り分け）・スコープを一括で確定。クライアント公開資料の3構成から採用案を確定。",
         "要件定義書、確定ワイヤーフレーム、権限マトリクス",
         "Must範囲（改善対象14課題＋AI-1・AI-2）を承認",
         "Go／No-Go②：要件定義書・確定ワイヤーフレーム・権限マトリクスの承認。",
         "マーケティング責任者／SharePoint管理者・IT窓口"),
        ("F2　構築",
         "目安4〜8週間\n（規模により変動）",
         "SharePoint（標準Web パーツ）とAI（Copilot Studio／SharePoint agent）を並行して構築。",
         "動作するSharePointページ一式、AI-1・AI-2のエージェント試作",
         "改善対象14課題の実装完了、AI-1・AI-2が動作する",
         "Go／No-Go③：構築物の一次レビュー通過（品質保証の着手可否）。",
         "実装担当"),
        ("F3　全体検証・移行",
         "目安2〜3週間",
         "改善対象14課題の解消を検証。現状維持・対象外・良好・維持・要現物確認の8課題の状態確認、既存コンテンツの移行、権限・セキュリティ確認、AI-2の外部リンク提示を含むアクセス権範囲の実機検証。",
         "検証済みステージング環境、品質保証報告",
         "品質保証の完了と受入条件の充足",
         "Go／No-Go④：品質保証完了・受入条件の充足・最終承認。",
         "実装担当／情報セキュリティ"),
        ("F4　同日カットオーバー公開",
         "目安1日\n（公開作業。事前リハーサルは別途）",
         "全ページとAI-1・AI-2を1つのリリースとして一斉に公開。",
         "本番公開ポータル",
         "全ページとAI-1・AI-2が同時に稼働開始",
         "（F3のGo／No-Go④通過後に実施。以降の追加ゲートなし）",
         "マーケティング責任者／SharePoint管理者／実装担当"),
    ],
    heights=[104] * 5, bold_cols=(1,), body_size=9.5, header_size=9, header_h=30,
)


# ---------------------------------------------------------------------------
# シート12: KPI
# ---------------------------------------------------------------------------
ws = wb.create_sheet("KPI")
ws.sheet_view.showGridLines = False
LC = 17
set_col_widths(ws, 1, LC, width=13)
r = page_header_bar(ws, 1, 1, LC, "KPI（取得主体別の3層）", 12)
r += 1
r = note_box(
    ws, r, 1, LC,
    "外部のアクセス解析ツールは利用できないため、SharePoint標準の「サイト利用状況（Site usage）」と"
    "「検索利用状況レポート（Search usage reports）」で測定します。"
    "検索利用状況レポートはサイトコレクション管理者のみが閲覧できるため、サイト所有者が自分で取得できるサイト利用状況を主軸に設計しています。"
    "AI-1とAI-2は取得可否が大きく異なるため、指標を分けています（#7・#8）。",
    height=46, italic=True,
)
r += 1

kpi_spans = [(1, 1), (2, 5), (6, 9), (10, 12), (13, 15), (16, 17)]
kpi_headers = ["#", "指標", "測定方法", "評価時点", "目標設定の方法", "取得者"]
EVAL = "ベースライン（公開前）を取得→公開後30日→公開後90日で追跡し、推移を比較する。"


def kpi_group(r, title, fill, rows):
    r = section_title(ws, r, 1, LC, title, fill=fill, color=WHITE, size=12, height=22)
    r = table(ws, r, kpi_spans, kpi_headers, rows, heights=[68] * len(rows),
              header_fill=NAVY_80, header_size=9, header_h=26,
              center_cols=(1, 16), bold_cols=(1,), body_size=9.5)
    return r + 1


r = kpi_group(r, "A. サイト所有者が自分で取得できる指標（サイト利用状況・主軸）", NAVY, [
    ("1", "主要ページの閲覧数・一意閲覧者数（ブランドライブラリ、キャンディデート向け 等）",
     "サイト利用状況のページ別閲覧数・一意閲覧者数を、刷新の前後で比較する。",
     EVAL, "ベースライン比で＋X％等の相対目標を関係者と協議のうえ設定する（絶対値の事前目標は置かない）。",
     "サイト所有者"),
    ("2", "人気コンテンツ（どの資料・ページが実際に使われているか）",
     "サイト利用状況の「人気コンテンツ（Popular content）」。刷新で導線を張った資料の閲覧数増加を確認する。",
     EVAL, "導線を張った資料が人気コンテンツ上位に入ることを目標とする（順位ベース。絶対数値目標は置かない）。",
     "サイト所有者"),
    ("3", "平均滞在時間（探索コストの目安）",
     "サイト利用状況の平均滞在時間。滞在時間の短縮が必ずしも改善とは限らないため"
     "（探しやすくなって短時間で用が済む場合と、探せずに離脱する場合の両方があり得るため）、他指標と併せて解釈する。",
     EVAL, "単独の数値目標は設定せず、他指標（#1・#2・#6）とセットで傾向を評価する。",
     "サイト所有者"),
    ("4", "サイト全体の訪問数・利用者数の推移",
     "サイト利用状況の「Site visits／Unique viewers／Popular Platforms（デバイス傾向）」。",
     EVAL, "ベースライン比での増加傾向を確認する（絶対値の事前目標は置かない）。",
     "サイト所有者"),
])
r = kpi_group(r, "B. 管理者への依頼が必要な指標（検索利用状況レポート）", NAVY_80, [
    ("5", "ゼロ件検索・人気検索語（「見つからない」の直接的な証拠）",
     "検索利用状況レポート（Site settings ＞ Site collection administration ＞ Microsoft Search ＞ Insights）。"
     "月次でサイトコレクション管理者にExcelエクスポートを依頼する運用とする（権限付与が可能ならそれが最善）。",
     EVAL, "ベースライン比でのゼロ件検索率の減少を目標とする。",
     "サイトコレクション管理者\n（依頼ベース）"),
])
r = kpi_group(r, "C. 権限が不要な手動指標・AIの利用状況（補完）", NAVY_40, [
    ("6", "マーケ部への「資料が見つからない」問い合わせ件数",
     "手動集計（現状値を取得し、刷新後の減少を確認する）。",
     EVAL, "ベースライン比での件数減少を目標とする。", "マーケ部（手動）"),
    ("7", "AI-1の利用状況・ブランドレビュー申請数",
     "Copilot Studioの分析機能、または Lists／Forms 等の標準集計機能で件数を取得する（実装する仕組みの設計に依存する）。",
     EVAL, "公開後30日の実績を踏まえてF3以降に目標値を設定する（事前の絶対値目標は置かない）。",
     "マーケ部（Lists／Forms 等の集計）"),
    ("8", "AI-2の利用状況",
     "取得方法・権限・保持期間は要確認。SharePoint agentの利用ログ取得機能はMicrosoft公式ドキュメントで確認できていないため、F3の実機検証で確認する。",
     "F3の実機検証で確認方法を確定したうえで、以降は他指標と同じ時点で追跡する。",
     "測定方法の確定後に設定する（現時点では未確定）。",
     "要確認（IT窓口経由の可能性）"),
])

r = note_box(
    ws, r, 1, LC,
    "運用：刷新前にベースラインを取得→目標値を設定→公開後30日・90日で測定します。"
    "検索利用状況レポートは過去31日（日次）／12か月（月次）でExcel出力できます。"
    "AI-2の利用状況（#8）は取得可否そのものが未確認のため、F3の実機検証の結果に応じて本表を更新します。",
    height=38, italic=True,
)


# ---------------------------------------------------------------------------
# シート13: AI-1・AI-2 実装と確認事項
# ---------------------------------------------------------------------------
ws = wb.create_sheet("AI-1・AI-2 実装と確認事項")
ws.sheet_view.showGridLines = False
LC = 14
set_col_widths(ws, 1, LC, width=13)
r = page_header_bar(ws, 1, 1, LC, "AI-1・AI-2 実装方式と確認事項", 13)
r += 1
r = note_box(ws, r, 1, LC,
             "Microsoft公式ドキュメントで確認できた事項と、実機検証が必要な事項を分けて記載しています。",
             height=24, italic=True)
r += 1

r = section_title(ws, r, 1, LC, "実装方式")
r = table(
    ws, r,
    spans=[(1, 2), (3, 8), (9, 14)],
    headers=["項目", "AI-1　ブランドセルフチェッカー", "AI-2　マーケ特化AIアシスタント"],
    rows=[
        ("実装方式",
         "Copilot Studioのカスタムエージェント（会話内のファイルアップロードと画像入力の分析）。SharePointページには埋め込み Web パーツでカスタムWebサイトチャネルを設置する。",
         "SharePointのエージェント機能（SharePoint agent）。サイト所有者の権限で作成し、メインエージェントに設定すると、ヘッダーのエージェントアイコンから全員が利用できる。"),
        ("配置場所",
         "「ブランドレビュー依頼」「ロゴ掲載申請」の文脈（トップの業務依頼セクション、ブランドライブラリのブランド使用・ロゴ掲載申請セクション）に埋め込み Web パーツで設置する。",
         "サイト全体のメインエージェントとしてトップページに設定する。全ページのヘッダーから利用できる。"),
        ("サイト所有者でできること",
         "Copilot Studioでのエージェント試作・トピック設計・ナレッジ登録・ファイルアップロード機能の有効化（テナントにCopilot Studioの契約と作成者権限がある場合）。",
         "SharePoint agentの作成、ナレッジの選定（対象ライブラリ・ページ）、メインエージェント設定、共有の制御。ライセンスが揃えば所有者だけで作成できる。"),
        ("管理者に確認すべきこと",
         "①M365 Copilot／Copilot Studioライセンスの契約・割当　②Copilot Studio Authorsロールの付与　③DLP／機密ラベルの参照除外設定　④カスタムWebサイトチャネル公開時の「認証なし」設定の許容可否　⑤PowerPointファイルの直接アップロード（実験的機能）の申請要否",
         "①M365 Copilotライセンスの契約・割当　②DLP／機密ラベルの参照除外設定　③SharePoint Advanced Managementの利用可否　④Graph connectors（外部データ取込）の要否"),
        ("回答範囲（実機検証が必要な事項）",
         "－（AI-1は画像入力とファイル解析が主用途のため対象外）",
         "ページ本文を根拠に回答できることはMicrosoft公式ドキュメントで確認済み。一方、ページ本文中の外部URLを期待どおり提示できるかは公式ドキュメントに記載がなく、F3で必須の実機検証項目とする。"),
        ("限界・注意点",
         "色のΔE数値照合やロゴのピクセル単位の検出といった厳密な判定は標準機能では保証されない。一次スクリーニング（気づきの提示）用途が現実的で、最終承認は人が行う運用を推奨する。PowerPointファイルの直接アップロードは実験的機能のため、当面はPDF変換を前段に挟む運用が現実的。",
         "回答は必ず利用者本人のアクセス権の範囲でセキュリティトリミングされる（権限境界は標準装備）。参照元への出典リンクが標準で付与される。"),
    ],
    heights=[96] * 6, bold_cols=(1,), center_cols=(1,),
)

r += 1
r = note_box(
    ws, r, 1, LC,
    "重要な制約：AI-2（SharePoint agent）は現行仕様でリストのデータを使いません"
    "（Microsoft公式ドキュメントに \"Agents currently don't use data from Lists.\" と明記）。"
    "クライアント公開資料などAI-2に拾わせたい情報は、リストではなくページ本文（テキスト Web パーツ）に記載する必要があります"
    "（詳細は「クライアント公開資料 構成比較」シート）。",
    height=46, fill="FDEBEC", bold=True,
)
r += 1
r = note_box(
    ws, r, 1, LC,
    "AI-1とAI-2は別のエージェントとして分離するのが妥当です。「文書検索の質問応答」と「ファイル解析＋ルール照合」で用途が大きく異なり、"
    "1つのエージェントに統合すると精度が落ちるためです。",
    height=28, italic=True,
)
r += 1
r = note_box(
    ws, r, 1, LC,
    "実機検証で確認する項目：①AI-2がページ本文中の外部URLを提示できるか　②AI-2の利用ログの取得可否　"
    "③AI-1のカスタムWebサイトチャネルの公開設定　④PowerPointファイルの直接アップロードの可否。"
    "いずれもF3で検証し、結果に応じて設計・運用を調整します。",
    height=34, fill="FFF4E0",
)


# ---------------------------------------------------------------------------
# シート14: リスク登録簿
# ---------------------------------------------------------------------------
ws = wb.create_sheet("リスク登録簿")
ws.sheet_view.showGridLines = False
LC = 17
set_col_widths(ws, 1, LC, width=13)
r = page_header_bar(ws, 1, 1, LC, "リスク登録簿（意思決定に必要な主要リスク）", 14)
r += 1
r = note_box(ws, r, 1, LC,
             "実装可否そのものに関わるリスク、および同日カットオーバー公開の前提を左右するリスクを優先して掲載しています。",
             height=24, italic=True)
r += 1
r = table(
    ws, r,
    spans=[(1, 4), (5, 8), (9, 13), (14, 17)],
    headers=["リスク", "影響", "対応策", "確認先"],
    rows=[
        ("M365 Copilot／Copilot Studio のライセンス・予算が確保できない",
         "AI-1・AI-2の実装可否と費用の前提が崩れ、同日カットオーバー公開そのものが成立しなくなる（最重要リスク）。",
         "F0を最優先で実施し、可否と費用を早期に確定する。確保できない場合は公開延期を原則とする（Must要件を外した部分公開は行わない）。",
         "マーケティング責任者／SharePoint管理者・IT窓口"),
        ("AI-1のカスタムWebサイトチャネル埋め込みで「認証なし」設定が必要になる可能性",
         "「認証なし」公開が社内ポリシー・情報セキュリティ基準に抵触する場合、この方式でAI-1を実装できず、代替実装の検討または見送りが必要になる。",
         "F0でIT窓口・情報セキュリティに公開設定の許容可否を確認する。許容できない場合はCopilot Studio内の別チャネル（Teams等）や他の埋め込み方式を代替として検討する。",
         "SharePoint管理者・IT窓口／情報セキュリティ"),
        ("AI-1のPowerPointファイル直接アップロードが実験的機能である",
         "実験的機能は予告なく仕様変更・提供停止される可能性があり、運用が不安定になる。",
         "当面はPDF変換を前段に挟む運用を既定とし、実験的機能への依存を避ける。機能の正式提供状況をF2〜F3で再確認する。",
         "実装担当"),
        ("AI-2の外部リンク提示能力がMicrosoft公式ドキュメントで確認できていない",
         "ページ本文中の外部URLを期待どおり提示できるかが不明で、クライアント向けの外部リンク案内が期待どおり機能しない可能性がある。",
         "F3で実機検証を必須タスク化し、期待どおり動作しない場合はページ本文の記述方法（アンカーテキストの明示化等）を調整する。",
         "実装担当（F3実機検証）"),
        ("AI-1のブランド準拠チェックの精度は一次スクリーニング止まりである",
         "色のΔE数値照合やロゴのピクセル単位の検出といった厳密な自動判定は標準機能で保証されないため、AIの判定のみではブランド逸脱を見逃す可能性がある。",
         "AI-1は「気づきの提示」用途と位置づけ、最終承認は人（ブランドレビュー担当）が行う運用を維持し、品質保証のプロセスを残す。",
         "マーケティング責任者"),
    ],
    heights=[82] * 5, bold_cols=(1,), body_size=9.5,
)


# ---------------------------------------------------------------------------
# 印刷設定
# ---------------------------------------------------------------------------
for _ws in wb.worksheets:
    _ws.page_setup.orientation = "landscape"
    _ws.page_setup.fitToWidth = 1
    _ws.page_setup.fitToHeight = 0
    _ws.sheet_properties.pageSetUpPr = WorksheetProperties(
        pageSetUpPr=PageSetupProperties(fitToPage=True)).pageSetUpPr
    _ws.page_margins.left = 0.3
    _ws.page_margins.right = 0.3
    _ws.page_margins.top = 0.4
    _ws.page_margins.bottom = 0.4

wb.save(OUT_PATH)
print("saved", OUT_PATH, round(os.path.getsize(OUT_PATH) / 1024 / 1024, 2), "MB")
print()
print("--- WFシートのコメント行と画像セクションの位置ズレ ---")
for rep in ALIGN_REPORT:
    print(f"{rep['sheet']}: sections={rep['sections']} scale={rep['scale']} "
          f"表示高={rep['disp_h']}px rows={rep['rows_used']} 最大ズレ={rep['max_drift_px']}px")
    for no, target, actual, d in rep["detail"]:
        print(f"    No.{no:>2}  理想={target:>7.1f}px  実配置={actual:>7.1f}px  ズレ={d:+.2f}px")
